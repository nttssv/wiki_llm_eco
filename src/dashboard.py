"""Read-only Streamlit dashboard for the narrative database."""

from __future__ import annotations

from html import escape
import json
from pathlib import Path
import sys
from typing import Any
from zipfile import ZipFile

import altair as alt
from docx import Document
import networkx as nx
import pandas as pd
from pyvis.network import Network
import streamlit as st
import streamlit.components.v1 as components

try:
    from .date_utils import week_label_for_published_date
    from .graph_rag import answer_question
    from .config import get_neo4j_database
    from .neo4j_store import neo4j_driver
    from .paths import EXPORTS_DIR, PROCESSED_MANIFEST_PATH, PROJECT_ROOT
    from .processed_registry import processed_article_records
except ImportError:
    PROJECT_ROOT = Path(__file__).resolve().parents[1]
    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    from src.date_utils import week_label_for_published_date
    from src.graph_rag import answer_question
    from src.config import get_neo4j_database
    from src.neo4j_store import neo4j_driver
    from src.paths import EXPORTS_DIR, PROCESSED_MANIFEST_PATH, PROJECT_ROOT
    from src.processed_registry import processed_article_records


GRAPH_EXPORT_PATH = EXPORTS_DIR / "graph.json"
GRAPH_LATEST_PATH = EXPORTS_DIR / "graph_latest.json"
DOCX_PREVIEW_LIMIT = 12000
DOCX_IMAGE_LIMIT = 4
NODE_COLORS = {
    "article": "#1d4ed8",
    "entity": "#0f766e",
    "theme": "#c2410c",
    "narrative": "#7c3f00",
}
NODE_SHAPES = {
    "article": "box",
    "entity": "dot",
    "theme": "triangle",
    "narrative": "diamond",
}

STATUS_BADGE_CLASSES: dict[str, str] = {
    "emerging": "status-emerging",
    "strengthening": "status-strengthening",
    "weakening": "status-weakening",
    "stable": "status-stable",
    "reversed": "status-reversed",
    "transitioning": "status-transitioning",
}

TREND_BADGE_CLASSES: dict[str, str] = {
    "new": "trend-new",
    "strengthening": "trend-strengthening",
    "recurring": "trend-recurring",
    "weakening": "trend-weakening",
}

ENTITY_TYPE_CLASSES: dict[str, str] = {
    "person": "etype-person",
    "company": "etype-company",
    "organization": "etype-organization",
    "country": "etype-country",
    "technology": "etype-technology",
}


def _path_version(path: Path) -> tuple[int, int]:
    """Return a stable cache key component for a file-backed data source."""

    if not path.exists():
        return (0, 0)

    stat = path.stat()
    return (stat.st_mtime_ns, stat.st_size)


def _empty_narratives_frame() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "id",
            "name",
            "thesis",
            "status",
            "importance_score",
            "first_seen_date",
            "last_seen_date",
            "mention_count",
            "linked_articles",
            "article_ids",
        ]
    )


def _empty_entities_frame() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "id",
            "name",
            "type",
            "description",
            "linked_articles",
            "article_ids",
            "graph_relationships",
        ]
    )


@st.cache_data(show_spinner=False)
def load_metrics(registry_version: tuple[int, int]) -> dict[str, int]:
    """Load top-line metrics from processed extraction artifacts."""

    del registry_version
    return {
        "articles": len(processed_article_records()),
        "entities": 0,
        "themes": 0,
        "narratives": 0,
        "graph_edges": 0,
    }


@st.cache_data(show_spinner=False)
def load_articles(registry_version: tuple[int, int]) -> pd.DataFrame:
    """Load article-level dashboard data from processed extraction artifacts."""

    del registry_version
    columns = [
        "id",
        "title",
        "source",
        "published_date",
        "category",
        "summary",
        "importance_score",
        "original_file_path",
        "linked_entities",
        "linked_themes",
        "linked_narratives",
        "week_label",
        "selection_label",
    ]
    records = []
    for record in processed_article_records():
        records.append(
            {
                **record,
                "linked_entities": "",
                "linked_themes": "",
                "linked_narratives": "",
            }
        )
    dataframe = pd.DataFrame(records)
    if dataframe.empty:
        return pd.DataFrame(columns=columns)

    dataframe["week_label"] = dataframe["published_date"].apply(
        lambda value: week_label_for_published_date(str(value or ""))
    )
    dataframe["selection_label"] = dataframe.apply(
        lambda row: f"{row['title']} | {row['published_date'] or 'Undated'} | {row['source']}",
        axis=1,
    )
    return dataframe[columns]


@st.cache_data(show_spinner=False)
def load_narratives(registry_version: tuple[int, int]) -> pd.DataFrame:
    """Return an empty narrative frame when Neo4j graph data is unavailable."""

    del registry_version
    return _empty_narratives_frame()


@st.cache_data(show_spinner=False)
def load_narrative_trends(week_label: str, registry_version: tuple[int, int]) -> pd.DataFrame:
    """Return an empty trend frame when Neo4j graph data is unavailable."""

    del week_label, registry_version
    return _empty_narrative_trends_frame()


def _empty_narrative_trends_frame() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "id",
            "name",
            "thesis",
            "extracted_status",
            "trend_status",
            "importance_score",
            "mention_count",
            "first_seen_date",
            "last_seen_date",
            "week_mentions",
        ],
    )


def _classify_narrative_trend(week_label: str, first_seen_date: str, last_seen_date: str, week_mentions: int) -> str:
    if first_seen_date == week_label and week_mentions > 0:
        return "NEW"
    if week_mentions > 1:
        return "STRENGTHENING"
    if week_mentions == 1 and first_seen_date and first_seen_date < week_label:
        return "RECURRING"
    if week_mentions == 0 and first_seen_date and first_seen_date < week_label and last_seen_date < week_label:
        return "WEAKENING"
    return "NEW" if week_mentions > 0 else "WEAKENING"


@st.cache_data(show_spinner=False, ttl=30)
def load_narrative_trends_neo4j(week_label: str, cache_token: int) -> pd.DataFrame:
    """Load narrative trend classifications from Neo4j."""

    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            rows = session.run(
                """
                MATCH (n:Narrative)
                OPTIONAL MATCH (a:Article)-[:HAS_NARRATIVE]->(n)
                WITH n, [article IN collect(DISTINCT a) WHERE article IS NOT NULL] AS articles
                WITH n, articles, [article IN articles WHERE article.week = $week] AS week_articles
                RETURN
                    n.id AS id,
                    coalesce(n.name, '') AS name,
                    coalesce(n.thesis, '') AS thesis,
                    coalesce(n.status, '') AS extracted_status,
                    coalesce(n.importance_score, 0) AS importance_score,
                    coalesce(n.mention_count, size(articles)) AS mention_count,
                    coalesce(n.first_seen_date, '') AS first_seen_date,
                    coalesce(n.last_seen_date, '') AS last_seen_date,
                    size(week_articles) AS week_mentions
                ORDER BY mention_count DESC, importance_score DESC, name ASC
                """,
                week=week_label,
            ).data()

    if not rows:
        return _empty_narrative_trends_frame()

    trend_rows: list[dict[str, Any]] = []
    for row in rows:
        first_seen_date = str(row.get("first_seen_date") or "")
        last_seen_date = str(row.get("last_seen_date") or "")
        week_mentions = int(row.get("week_mentions") or 0)
        if first_seen_date and first_seen_date > week_label and week_mentions == 0:
            continue
        trend_rows.append(
            {
                "id": str(row.get("id") or ""),
                "name": str(row.get("name") or ""),
                "thesis": str(row.get("thesis") or ""),
                "extracted_status": str(row.get("extracted_status") or ""),
                "trend_status": _classify_narrative_trend(
                    week_label,
                    first_seen_date,
                    last_seen_date,
                    week_mentions,
                ),
                "importance_score": int(row.get("importance_score") or 0),
                "mention_count": int(row.get("mention_count") or 0),
                "first_seen_date": first_seen_date,
                "last_seen_date": last_seen_date,
                "week_mentions": week_mentions,
            }
        )

    if not trend_rows:
        return _empty_narrative_trends_frame()
    return pd.DataFrame(trend_rows, columns=_empty_narrative_trends_frame().columns)


@st.cache_data(show_spinner=False)
def load_entities(registry_version: tuple[int, int]) -> pd.DataFrame:
    """Return an empty entity frame when Neo4j graph data is unavailable."""

    del registry_version
    return _empty_entities_frame()


def _join_values(values: Any, separator: str) -> str:
    if not isinstance(values, list):
        return ""
    return separator.join(str(value) for value in values if value)


def _first_sentence(text: str, limit: int = 220) -> str:
    compact = " ".join(str(text or "").split())
    if not compact:
        return ""
    sentence_end = compact.find(". ")
    if 0 <= sentence_end <= limit:
        return compact[: sentence_end + 1]
    return compact[:limit].rstrip()


@st.cache_data(show_spinner=False, ttl=30)
def neo4j_dashboard_available(cache_token: int = 0) -> bool:
    """Return whether the dashboard can read the Neo4j graph."""

    try:
        with neo4j_driver() as driver:
            driver.verify_connectivity()
        return True
    except Exception:
        return False


@st.cache_data(show_spinner=False, ttl=30)
def load_metrics_neo4j(cache_token: int) -> dict[str, int]:
    """Load top-line dashboard metrics from Neo4j."""

    statements = {
        "articles": "MATCH (n:Article) RETURN count(n) AS count",
        "entities": "MATCH (n:Entity) RETURN count(n) AS count",
        "themes": "MATCH (n:Theme) RETURN count(n) AS count",
        "narratives": "MATCH (n:Narrative) RETURN count(n) AS count",
        "graph_edges": "MATCH ()-[r:RELATES_TO]->() RETURN count(r) AS count",
    }
    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            metrics: dict[str, int] = {}
            for key, statement in statements.items():
                row = session.run(statement).single()
                metrics[key] = int(row["count"] if row else 0)
    return metrics


@st.cache_data(show_spinner=False, ttl=30)
def load_articles_neo4j(cache_token: int) -> pd.DataFrame:
    """Load article-level dashboard data from Neo4j."""

    columns = [
        "id",
        "title",
        "source",
        "published_date",
        "category",
        "summary",
        "importance_score",
        "original_file_path",
        "linked_entities",
        "linked_themes",
        "linked_narratives",
        "week_label",
        "selection_label",
    ]
    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            rows = session.run(
                """
                MATCH (a:Article)
                OPTIONAL MATCH (a)-[:IN_WEEK]->(w:Week)
                WITH a, w
                OPTIONAL MATCH (a)-[:MENTIONS]->(e:Entity)
                WITH a, w, collect(DISTINCT e.name) AS entities
                OPTIONAL MATCH (a)-[:HAS_THEME]->(t:Theme)
                WITH a, w, entities, collect(DISTINCT t.name) AS themes
                OPTIONAL MATCH (a)-[:HAS_NARRATIVE]->(n:Narrative)
                WITH a, w, entities, themes, collect(DISTINCT n.name) AS narratives
                RETURN
                    a.id AS id,
                    coalesce(a.title, '') AS title,
                    coalesce(a.source, '') AS source,
                    coalesce(a.published_date, '') AS published_date,
                    coalesce(a.category, '') AS category,
                    coalesce(a.summary, '') AS summary,
                    coalesce(a.importance_score, 0) AS importance_score,
                    coalesce(a.original_file_path, '') AS original_file_path,
                    coalesce(a.week, w.label, '') AS week_label,
                    entities,
                    themes,
                    narratives
                ORDER BY published_date DESC, title ASC
                """
            ).data()

    dataframe = pd.DataFrame(rows)
    if dataframe.empty:
        return pd.DataFrame(columns=columns)

    dataframe["linked_entities"] = dataframe["entities"].apply(lambda values: _join_values(values, ", "))
    dataframe["linked_themes"] = dataframe["themes"].apply(lambda values: _join_values(values, ", "))
    dataframe["linked_narratives"] = dataframe["narratives"].apply(lambda values: _join_values(values, ", "))
    dataframe["week_label"] = dataframe.apply(
        lambda row: row["week_label"] or week_label_for_published_date(str(row["published_date"] or "")),
        axis=1,
    )
    dataframe["selection_label"] = dataframe.apply(
        lambda row: f"{row['title']} | {row['published_date'] or 'Undated'} | {row['source']}",
        axis=1,
    )
    return dataframe[columns]


@st.cache_data(show_spinner=False, ttl=30)
def load_narratives_neo4j(cache_token: int) -> pd.DataFrame:
    """Load narrative-level dashboard data from Neo4j."""

    columns = [
        "id",
        "name",
        "thesis",
        "status",
        "importance_score",
        "first_seen_date",
        "last_seen_date",
        "mention_count",
        "linked_articles",
        "article_ids",
    ]
    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            rows = session.run(
                """
                MATCH (n:Narrative)
                OPTIONAL MATCH (a:Article)-[:HAS_NARRATIVE]->(n)
                WITH n, collect(DISTINCT a.title) AS article_titles, collect(DISTINCT a.id) AS article_ids
                RETURN
                    n.id AS id,
                    coalesce(n.name, '') AS name,
                    coalesce(n.thesis, '') AS thesis,
                    coalesce(n.status, '') AS status,
                    coalesce(n.importance_score, 0) AS importance_score,
                    coalesce(n.first_seen_date, '') AS first_seen_date,
                    coalesce(n.last_seen_date, '') AS last_seen_date,
                    coalesce(n.mention_count, size(article_ids)) AS mention_count,
                    article_titles,
                    article_ids
                ORDER BY mention_count DESC, importance_score DESC, name ASC
                """
            ).data()

    dataframe = pd.DataFrame(rows)
    if dataframe.empty:
        return pd.DataFrame(columns=columns)

    dataframe["linked_articles"] = dataframe["article_titles"].apply(lambda values: _join_values(values, " | "))
    dataframe["article_ids"] = dataframe["article_ids"].apply(lambda values: _join_values(values, "|"))
    return dataframe[columns]


@st.cache_data(show_spinner=False, ttl=30)
def load_entities_neo4j(cache_token: int) -> pd.DataFrame:
    """Load entity-level dashboard data from Neo4j."""

    columns = [
        "id",
        "name",
        "type",
        "description",
        "linked_articles",
        "article_ids",
        "graph_relationships",
    ]
    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            rows = session.run(
                """
                MATCH (e:Entity)
                OPTIONAL MATCH (a:Article)-[:MENTIONS]->(e)
                WITH e, collect(DISTINCT a.title) AS article_titles, collect(DISTINCT a.id) AS article_ids
                OPTIONAL MATCH (e)-[out:RELATES_TO]->(target:Entity)
                WITH e, article_titles, article_ids,
                     collect(DISTINCT e.name + ' --' + coalesce(out.relationship, 'RELATES_TO') + '--> ' + target.name) AS outgoing
                OPTIONAL MATCH (source:Entity)-[inc:RELATES_TO]->(e)
                WITH e, article_titles, article_ids, outgoing,
                     collect(DISTINCT source.name + ' --' + coalesce(inc.relationship, 'RELATES_TO') + '--> ' + e.name) AS incoming
                RETURN
                    e.id AS id,
                    coalesce(e.name, '') AS name,
                    coalesce(e.type, '') AS type,
                    coalesce(e.description, '') AS description,
                    article_titles,
                    article_ids,
                    outgoing + incoming AS relationships
                ORDER BY name ASC
                """
            ).data()

    dataframe = pd.DataFrame(rows)
    if dataframe.empty:
        return pd.DataFrame(columns=columns)

    dataframe["linked_articles"] = dataframe["article_titles"].apply(lambda values: _join_values(values, " | "))
    dataframe["article_ids"] = dataframe["article_ids"].apply(lambda values: _join_values(values, "|"))
    dataframe["graph_relationships"] = dataframe["relationships"].apply(lambda values: _join_values(values, " | "))
    return dataframe[columns]


def _graph_snapshot_path(week_label: str | None) -> Path:
    if not week_label or week_label == "latest":
        return GRAPH_LATEST_PATH if GRAPH_LATEST_PATH.exists() else GRAPH_EXPORT_PATH
    return EXPORTS_DIR / f"graph_{week_label}.json"


def list_available_graph_weeks() -> list[str]:
    return sorted(
        [
            path.stem.removeprefix("graph_")
            for path in EXPORTS_DIR.glob("graph_*.json")
            if path.stem not in {"graph_latest"}
        ]
    )


@st.cache_data(show_spinner=False, ttl=30)
def list_available_graph_weeks_neo4j(cache_token: int) -> list[str]:
    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            rows = session.run(
                """
                MATCH (a:Article)
                WHERE coalesce(a.week, '') <> ''
                RETURN DISTINCT a.week AS week
                ORDER BY week ASC
                """
            ).data()
    return [str(row["week"]) for row in rows if row.get("week")]


@st.cache_data(show_spinner=False)
def load_graph_snapshot(week_label: str | None, graph_version: tuple[int, int]) -> dict[str, Any] | None:
    """Load one exported graph snapshot."""

    graph_path = _graph_snapshot_path(week_label)
    if not graph_path.exists():
        return None
    return json.loads(graph_path.read_text(encoding="utf-8"))


@st.cache_data(show_spinner=False, ttl=30)
def load_graph_snapshot_neo4j(week_label: str | None, cache_token: int) -> dict[str, Any] | None:
    """Build a graph-view snapshot directly from Neo4j."""

    selected_week = week_label
    if not selected_week or selected_week == "latest":
        weeks = list_available_graph_weeks_neo4j(cache_token)
        selected_week = weeks[-1] if weeks else None
    if not selected_week:
        return None

    with neo4j_driver() as driver:
        with driver.session(database=get_neo4j_database()) as session:
            article_rows = session.run(
                """
                MATCH (a:Article)
                WHERE a.week = $week
                RETURN
                    a.id AS id,
                    coalesce(a.title, '') AS title,
                    coalesce(a.summary, '') AS summary,
                    coalesce(a.importance_score, 0) AS importance_score,
                    coalesce(a.published_date, '') AS published_date,
                    coalesce(a.source, '') AS source
                ORDER BY published_date DESC, title ASC
                """,
                week=selected_week,
            ).data()
            article_ids = [str(row["id"]) for row in article_rows if row.get("id")]

            context_rows = session.run(
                """
                MATCH (a:Article)-[r:MENTIONS|HAS_THEME|HAS_NARRATIVE]->(node)
                WHERE a.id IN $article_ids
                RETURN
                    a.id AS article_id,
                    coalesce(a.title, '') AS article_title,
                    coalesce(a.summary, '') AS article_summary,
                    type(r) AS relationship,
                    labels(node) AS labels,
                    node.id AS node_id,
                    coalesce(node.name, node.title, node.id) AS label,
                    coalesce(node.description, node.thesis, '') AS detail,
                    coalesce(node.importance_score, 0) AS importance_score,
                    coalesce(node.mention_count, 0) AS mention_count
                ORDER BY article_title ASC, label ASC
                """,
                article_ids=article_ids,
            ).data() if article_ids else []

            relationship_rows = session.run(
                """
                MATCH (source:Entity)-[rel:RELATES_TO]->(target:Entity)
                MATCH (article:Article {id: rel.evidence_article_id})
                WHERE article.id IN $article_ids
                RETURN
                    source.id AS source_id,
                    coalesce(source.name, '') AS source_label,
                    coalesce(source.description, '') AS source_detail,
                    coalesce(source.type, '') AS source_entity_type,
                    target.id AS target_id,
                    coalesce(target.name, '') AS target_label,
                    coalesce(target.description, '') AS target_detail,
                    coalesce(target.type, '') AS target_entity_type,
                    coalesce(rel.relationship, 'RELATES_TO') AS relationship,
                    coalesce(rel.confidence, 0.0) AS confidence,
                    article.id AS article_id,
                    coalesce(article.title, '') AS article_title,
                    coalesce(article.summary, '') AS article_summary
                ORDER BY article.published_date DESC, confidence DESC
                """,
                article_ids=article_ids,
            ).data() if article_ids else []

    nodes_by_id: dict[str, dict[str, Any]] = {}
    linked_articles_by_node: dict[str, set[str]] = {}
    edges: list[dict[str, Any]] = []

    for row in article_rows:
        article_id = str(row.get("id") or "")
        if not article_id:
            continue
        nodes_by_id[article_id] = {
            "id": article_id,
            "label": str(row.get("title") or article_id),
            "type": "article",
            "detail": str(row.get("summary") or ""),
            "importance_score": int(row.get("importance_score") or 0),
            "linked_articles": 1,
            "mention_count": 0,
            "week": selected_week,
        }

    def add_related_node(
        node_id: str,
        label: str,
        node_type: str,
        detail: str,
        importance_score: int = 0,
        mention_count: int = 0,
    ) -> None:
        if not node_id:
            return
        existing = nodes_by_id.get(node_id)
        if existing:
            existing["linked_articles"] = len(linked_articles_by_node.get(node_id, set()))
            return
        nodes_by_id[node_id] = {
            "id": node_id,
            "label": label or node_id,
            "type": node_type,
            "detail": detail or "",
            "importance_score": int(importance_score or 0),
            "linked_articles": len(linked_articles_by_node.get(node_id, set())),
            "mention_count": int(mention_count or 0),
            "week": selected_week,
        }

    for row in context_rows:
        node_id = str(row.get("node_id") or "")
        labels = row.get("labels") if isinstance(row.get("labels"), list) else []
        node_type = "entity"
        if "Theme" in labels:
            node_type = "theme"
        elif "Narrative" in labels:
            node_type = "narrative"
        article_id = str(row.get("article_id") or "")
        if article_id and node_id:
            linked_articles_by_node.setdefault(node_id, set()).add(article_id)
        add_related_node(
            node_id,
            str(row.get("label") or ""),
            node_type,
            str(row.get("detail") or ""),
            int(row.get("importance_score") or 0),
            int(row.get("mention_count") or 0),
        )
        if article_id and node_id:
            edges.append(
                {
                    "source": article_id,
                    "target": node_id,
                    "relationship": str(row.get("relationship") or ""),
                    "evidence_article_id": article_id,
                    "confidence": 1.0,
                    "evidence_title": str(row.get("article_title") or ""),
                    "narrative_sentence": _first_sentence(str(row.get("article_summary") or "")),
                    "week": selected_week,
                }
            )

    for row in relationship_rows:
        article_id = str(row.get("article_id") or "")
        source_id = str(row.get("source_id") or "")
        target_id = str(row.get("target_id") or "")
        if article_id and source_id:
            linked_articles_by_node.setdefault(source_id, set()).add(article_id)
        if article_id and target_id:
            linked_articles_by_node.setdefault(target_id, set()).add(article_id)
        add_related_node(
            source_id,
            str(row.get("source_label") or ""),
            "entity",
            str(row.get("source_detail") or ""),
        )
        add_related_node(
            target_id,
            str(row.get("target_label") or ""),
            "entity",
            str(row.get("target_detail") or ""),
        )
        if source_id and target_id:
            edges.append(
                {
                    "source": source_id,
                    "target": target_id,
                    "relationship": str(row.get("relationship") or "RELATES_TO"),
                    "evidence_article_id": article_id,
                    "confidence": float(row.get("confidence") or 0.0),
                    "evidence_title": str(row.get("article_title") or ""),
                    "narrative_sentence": (
                        f"{row.get('source_label') or source_id} --{row.get('relationship') or 'RELATES_TO'}--> "
                        f"{row.get('target_label') or target_id}"
                    ),
                    "week": selected_week,
                }
            )

    for node_id, article_links in linked_articles_by_node.items():
        if node_id in nodes_by_id:
            nodes_by_id[node_id]["linked_articles"] = len(article_links)

    return {
        "week": selected_week,
        "nodes": list(nodes_by_id.values()),
        "edges": edges,
    }


def _node_signature(node: dict[str, Any]) -> tuple[str, str]:
    return str(node.get("type", "")), str(node.get("label", ""))


def _edge_signature(edge: dict[str, Any], node_by_id: dict[str, dict[str, Any]]) -> tuple[str, str, str]:
    source_node = node_by_id.get(str(edge.get("source")), {})
    target_node = node_by_id.get(str(edge.get("target")), {})
    return (
        str(source_node.get("label", edge.get("source", ""))),
        str(edge.get("relationship", "")),
        str(target_node.get("label", edge.get("target", ""))),
    )


def build_comparison_payload(
    current_payload: dict[str, Any],
    previous_payload: dict[str, Any] | None,
) -> tuple[dict[str, Any], dict[str, int], list[tuple[str, int]]]:
    current_nodes = current_payload.get("nodes", [])
    current_edges = current_payload.get("edges", [])
    current_node_by_id = {str(node["id"]): node for node in current_nodes}
    current_node_signatures = {_node_signature(node): node for node in current_nodes}
    current_edge_signatures = {
        _edge_signature(edge, current_node_by_id): edge
        for edge in current_edges
    }

    previous_nodes = previous_payload.get("nodes", []) if previous_payload else []
    previous_edges = previous_payload.get("edges", []) if previous_payload else []
    previous_node_by_id = {str(node["id"]): node for node in previous_nodes}
    previous_node_signatures = {_node_signature(node): node for node in previous_nodes}
    previous_edge_signatures = {
        _edge_signature(edge, previous_node_by_id): edge
        for edge in previous_edges
    }

    new_node_signatures = set(current_node_signatures) - set(previous_node_signatures)
    removed_node_signatures = set(previous_node_signatures) - set(current_node_signatures)
    new_edge_signatures = set(current_edge_signatures) - set(previous_edge_signatures)
    removed_edge_signatures = set(previous_edge_signatures) - set(current_edge_signatures)

    merged_nodes: list[dict[str, Any]] = []
    signature_to_node_id: dict[tuple[str, str], str] = {}
    for signature, node in current_node_signatures.items():
        compare_status = "new" if signature in new_node_signatures else "unchanged"
        merged_node = {**node, "compare_status": compare_status}
        merged_nodes.append(merged_node)
        signature_to_node_id[signature] = str(node["id"])

    for signature in removed_node_signatures:
        node = previous_node_signatures[signature]
        removed_id = f"removed::{node['id']}"
        merged_node = {**node, "id": removed_id, "compare_status": "removed"}
        merged_nodes.append(merged_node)
        signature_to_node_id[signature] = removed_id

    merged_edges: list[dict[str, Any]] = []
    for signature, edge in current_edge_signatures.items():
        source_signature = _node_signature(current_node_by_id[str(edge["source"])])
        target_signature = _node_signature(current_node_by_id[str(edge["target"])])
        merged_edges.append(
            {
                **edge,
                "source": signature_to_node_id[source_signature],
                "target": signature_to_node_id[target_signature],
                "compare_status": "new" if signature in new_edge_signatures else "unchanged",
            }
        )

    for signature in removed_edge_signatures:
        edge = previous_edge_signatures[signature]
        source_signature = _node_signature(previous_node_by_id[str(edge["source"])])
        target_signature = _node_signature(previous_node_by_id[str(edge["target"])])
        merged_edges.append(
            {
                **edge,
                "source": signature_to_node_id[source_signature],
                "target": signature_to_node_id[target_signature],
                "compare_status": "removed",
            }
        )

    current_entity_degrees: dict[str, int] = {}
    previous_entity_degrees: dict[str, int] = {}
    for signature in current_edge_signatures:
        source_label, _, target_label = signature
        if ("entity", source_label) in current_node_signatures:
            current_entity_degrees[source_label] = current_entity_degrees.get(source_label, 0) + 1
        if ("entity", target_label) in current_node_signatures:
            current_entity_degrees[target_label] = current_entity_degrees.get(target_label, 0) + 1
    for signature in previous_edge_signatures:
        source_label, _, target_label = signature
        if ("entity", source_label) in previous_node_signatures:
            previous_entity_degrees[source_label] = previous_entity_degrees.get(source_label, 0) + 1
        if ("entity", target_label) in previous_node_signatures:
            previous_entity_degrees[target_label] = previous_entity_degrees.get(target_label, 0) + 1
    emerging_entities = sorted(
        [
            (label, current_entity_degrees.get(label, 0) - previous_entity_degrees.get(label, 0))
            for label in current_entity_degrees
            if current_entity_degrees.get(label, 0) - previous_entity_degrees.get(label, 0) > 0
        ],
        key=lambda item: item[1],
        reverse=True,
    )[:5]

    summary = {
        "current_nodes": len(current_nodes),
        "previous_nodes": len(previous_nodes),
        "new_nodes": len(new_node_signatures),
        "removed_nodes": len(removed_node_signatures),
        "new_edges": len(new_edge_signatures),
        "removed_edges": len(removed_edge_signatures),
    }

    merged_payload = {
        "week": current_payload.get("week", ""),
        "nodes": merged_nodes,
        "edges": merged_edges,
    }
    return merged_payload, summary, emerging_entities


def _split_ids(serialized_ids: str) -> set[str]:
    return {item for item in str(serialized_ids or "").split("|") if item}


def _split_names(serialized_names: str) -> list[str]:
    return [item.strip() for item in str(serialized_names or "").split(",") if item.strip()]


def _resolve_article_file_path(file_path: str) -> Path | None:
    raw_value = str(file_path or "").strip()
    if not raw_value:
        return None

    candidate = Path(raw_value)
    if candidate.is_absolute() and candidate.exists():
        return candidate

    if not candidate.is_absolute():
        project_relative = PROJECT_ROOT / candidate
        if project_relative.exists():
            return project_relative

    parts = candidate.parts
    if "data" in parts:
        data_index = parts.index("data")
        remapped = PROJECT_ROOT.joinpath(*parts[data_index:])
        if remapped.exists():
            return remapped

    return None


def _is_docx_preview_truncated(file_path: str) -> bool:
    resolved_path = _resolve_article_file_path(file_path)
    if resolved_path is None:
        return False

    try:
        paragraphs = [
            paragraph.text.strip()
            for paragraph in Document(resolved_path).paragraphs
            if paragraph.text.strip()
        ]
    except Exception:
        return False

    preview_text = "\n\n".join(paragraphs)
    return len(preview_text) > DOCX_PREVIEW_LIMIT


@st.cache_data(show_spinner=False)
def preview_docx(file_path: str) -> str:
    """Render a live preview from the original DOCX file."""

    resolved_path = _resolve_article_file_path(file_path)
    if resolved_path is None:
        return "Original DOCX not found"

    try:
        doc = Document(resolved_path)
    except Exception as e:
        return f"Error loading DOCX: {e}"

    paragraphs = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    ]
    preview_text = "\n\n".join(paragraphs)
    return preview_text[:DOCX_PREVIEW_LIMIT]


@st.cache_data(show_spinner=False)
def preview_docx_images(file_path: str) -> list[bytes]:
    """Extract embedded images from the original DOCX for dashboard display."""

    resolved_path = _resolve_article_file_path(file_path)
    if resolved_path is None:
        return []

    image_suffixes = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    image_payloads: list[bytes] = []

    try:
        with ZipFile(resolved_path) as archive:
            media_names = sorted(
                name
                for name in archive.namelist()
                if name.startswith("word/media/")
                and Path(name).suffix.lower() in image_suffixes
            )
            for media_name in media_names[:DOCX_IMAGE_LIMIT]:
                image_payloads.append(archive.read(media_name))
    except Exception:
        # If ZIP is corrupt or not a ZIP, return empty list
        return []

    return image_payloads


def _filter_articles(
    articles_df: pd.DataFrame,
    week_filter: str,
    source_filter: str,
    category_filter: str,
    search_term: str,
) -> pd.DataFrame:
    filtered = articles_df.copy()

    if week_filter != "All":
        filtered = filtered[filtered["week_label"] == week_filter]

    if source_filter != "All":
        filtered = filtered[filtered["source"] == source_filter]

    if category_filter != "All":
        filtered = filtered[filtered["category"] == category_filter]

    if search_term:
        mask = (
            filtered["title"].str.contains(search_term, case=False, na=False)
            | filtered["summary"].str.contains(search_term, case=False, na=False)
            | filtered["linked_entities"].str.contains(search_term, case=False, na=False)
            | filtered["linked_narratives"].str.contains(search_term, case=False, na=False)
        )
        filtered = filtered[mask]

    return filtered


def _filter_related_rows(dataframe: pd.DataFrame, article_ids: set[str]) -> pd.DataFrame:
    if not article_ids or dataframe.empty or "article_ids" not in dataframe.columns:
        return dataframe.iloc[0:0].copy()
    mask = dataframe["article_ids"].apply(lambda value: bool(_split_ids(str(value)) & article_ids))
    return dataframe.loc[mask].copy()


def _node_visual_weight(attributes: dict[str, Any], visible_connections: int) -> tuple[int, str]:
    """Calculate a display size and importance label for graph nodes."""

    importance_score = int(attributes.get("importance_score", 0) or 0)
    linked_articles = int(attributes.get("linked_articles", 0) or 0)
    mention_count = int(attributes.get("mention_count", 0) or 0)

    if importance_score > 0:
        prominence = importance_score
        importance_label = f"Importance: {importance_score}/10"
    elif mention_count > 0:
        prominence = min(10, max(2, mention_count * 2))
        importance_label = f"Importance proxy: {mention_count} narrative mention(s)"
    else:
        prominence = min(10, max(2, linked_articles + visible_connections))
        importance_label = f"Importance proxy: {linked_articles} linked article(s)"

    size = 16 + (prominence * 2) + min(visible_connections, 6)
    return size, importance_label


def _looks_like_raw_graph_id(value: str) -> bool:
    lowered = str(value or "").strip().lower()
    return lowered.startswith(("entity_", "theme_", "narrative_", "article_"))


def _safe_canvas_label(label: str) -> str:
    clean_label = str(label or "").strip()
    if not clean_label or _looks_like_raw_graph_id(clean_label):
        return ""
    return clean_label


def render_graph(
    payload: dict[str, list[dict[str, Any]]],
    selected_types: list[str],
    selected_relationships: list[str],
    search_node: str,
) -> None:
    """Render a filtered graph using pyvis."""

    nodes = payload.get("nodes", [])
    edges = payload.get("edges", [])
    node_by_id = {str(node["id"]): node for node in nodes}

    article_scoped_edges = list(edges)
    edges = list(article_scoped_edges)

    if selected_relationships:
        filtered_edges: list[dict[str, Any]] = []
        for edge in edges:
            source_type = str(node_by_id.get(str(edge.get("source")), {}).get("type", ""))
            target_type = str(node_by_id.get(str(edge.get("target")), {}).get("type", ""))
            compare_status = str(edge.get("compare_status", ""))
            if compare_status == "removed":
                filtered_edges.append(edge)
                continue
            if source_type == "article" or target_type == "article":
                filtered_edges.append(edge)
                continue
            if str(edge.get("relationship", "")) in selected_relationships:
                filtered_edges.append(edge)
        edges = filtered_edges

    node_search = search_node.strip().lower()

    if selected_types and not node_search:
        allowed_type_ids = {
            str(node["id"])
            for node in nodes
            if str(node.get("type", "")) in selected_types
        }
        edges = [
            edge
            for edge in edges
            if str(edge.get("source")) in allowed_type_ids and str(edge.get("target")) in allowed_type_ids
        ]

    debug_selected_label = ""
    debug_neighbors: list[dict[str, str]] = []
    debug_edges: list[str] = []

    if node_search:
        matching_ids = {
            str(node["id"])
            for node in nodes
            if node_search in str(node.get("label", "")).lower()
        }
        neighborhood_seed_edges = [
            edge
            for edge in article_scoped_edges
            if str(edge.get("source")) in matching_ids or str(edge.get("target")) in matching_ids
        ]
        visible_node_ids = set(matching_ids)
        for edge in neighborhood_seed_edges:
            visible_node_ids.add(str(edge["source"]))
            visible_node_ids.add(str(edge["target"]))

        edges = [
            edge
            for edge in article_scoped_edges
            if str(edge.get("source")) in visible_node_ids and str(edge.get("target")) in visible_node_ids
        ]

        if matching_ids:
            selected_id = sorted(matching_ids)[0]
            debug_selected_label = str(node_by_id.get(selected_id, {}).get("label", ""))
            neighbor_ids = sorted(visible_node_ids - {selected_id})
            debug_neighbors = [
                {
                    "label": str(node_by_id.get(node_id, {}).get("label", node_id)),
                    "type": str(node_by_id.get(node_id, {}).get("type", "")),
                }
                for node_id in neighbor_ids
            ]
            debug_edges = [
                f"{node_by_id[str(edge['source'])]['label']} -- {edge['relationship']} --> {node_by_id[str(edge['target'])]['label']}"
                for edge in edges
            ]

    connected_ids = {str(edge["source"]) for edge in edges} | {str(edge["target"]) for edge in edges}
    show_visible_edge_labels = bool(node_search)
    graph = nx.DiGraph()

    for node_id in connected_ids:
        node = node_by_id.get(node_id)
        if node is None:
            continue
        node_type = str(node.get("type", "entity"))
        readable_label = _safe_canvas_label(str(node.get("label", "")))
        graph.add_node(
            node_id,
            label=readable_label,
            node_type=node_type,
            color=NODE_COLORS.get(node_type, "#6e7781"),
            compare_status=str(node.get("compare_status", "unchanged")),
            detail=str(node.get("detail", "")),
            importance_score=int(node.get("importance_score", 0) or 0),
            linked_articles=int(node.get("linked_articles", 0) or 0),
            mention_count=int(node.get("mention_count", 0) or 0),
            week=str(node.get("week", "")),
        )

    for edge in edges:
        source_id = str(edge["source"])
        target_id = str(edge["target"])
        if source_id not in graph.nodes or target_id not in graph.nodes:
            continue
        graph.add_edge(
            source_id,
            target_id,
            relationship=str(edge.get("relationship", "")),
            confidence=float(edge.get("confidence", 0.0)),
            narrative_sentence=str(edge.get("narrative_sentence", "")),
            evidence_title=str(edge.get("evidence_title", "")),
            compare_status=str(edge.get("compare_status", "unchanged")),
            week=str(edge.get("week", "")),
        )

    if graph.number_of_nodes() == 0:
        st.info("No graph data matches the current filters.")
        return

    if node_search:
        with st.expander("DEBUG SELECTED NODE"):
            st.write(f"Selected node: {debug_selected_label or search_node}")
            st.write("Neighbor labels and types:")
            if debug_neighbors:
                st.write(debug_neighbors)
            else:
                st.write("No neighbors")
            st.write("Edge relationships:")
            if debug_edges:
                st.write(debug_edges)
            else:
                st.write("No visible edges")

    network = Network(height="940px", width="100%", directed=True, bgcolor="#ffffff", font_color="#111111")
    network.barnes_hut()

    for node_id, attributes in graph.nodes(data=True):
        visible_connections = int(graph.degree(node_id))
        node_size, importance_label = _node_visual_weight(attributes, visible_connections)
        detail = str(attributes.get("detail", "") or "")
        raw_id = str(node_id)
        tooltip_label = str(attributes["label"] or raw_id)
        compare_status = str(attributes.get("compare_status", "unchanged"))
        tooltip_lines = [
            tooltip_label,
            f"Type: {str(attributes['node_type'])}",
            f"Raw ID: {raw_id}",
            f"Week: {str(attributes.get('week', ''))}",
            importance_label,
            f"Visible connections: {visible_connections}",
        ]
        if compare_status != "unchanged":
            tooltip_lines.append(f"Change: {compare_status}")
        if detail:
            tooltip_lines.append(detail)
        node_color = str(attributes["color"])
        if compare_status == "new":
            node_color = "#16a34a"
        elif compare_status == "removed":
            node_color = "#dc2626"
        network.add_node(
            node_id,
            label=str(attributes["label"]),
            title="\n".join(tooltip_lines),
            color=node_color,
            shape=NODE_SHAPES.get(str(attributes["node_type"]), "dot"),
            size=node_size,
            borderWidth=max(2, node_size // 10),
            font={"size": 10, "face": "Georgia"},
        )

    for source_id, target_id, attributes in graph.edges(data=True):
        source_label = str(graph.nodes[source_id]["label"])
        target_label = str(graph.nodes[target_id]["label"])
        source_type = str(graph.nodes[source_id]["node_type"])
        target_type = str(graph.nodes[target_id]["node_type"])
        is_article_context_edge = source_type == "article" or target_type == "article"
        compare_status = str(attributes.get("compare_status", "unchanged"))
        narrative_sentence = str(attributes.get("narrative_sentence", "") or "")
        evidence_title = str(attributes.get("evidence_title", "") or "")
        tooltip_lines = [
            f"{source_label} → {target_label}",
            f"Relationship: {str(attributes['relationship'])}",
            f"Week: {str(attributes.get('week', ''))}",
        ]
        if compare_status != "unchanged":
            tooltip_lines.append(f"Change: {compare_status}")
        if narrative_sentence:
            tooltip_lines.append(f"Narrative: {narrative_sentence}")
        if evidence_title:
            tooltip_lines.append(f"Evidence article: {evidence_title}")
        tooltip_lines.append(f"Confidence: {attributes['confidence']:.2f}")
        edge_color = "#1d4ed8" if is_article_context_edge else "#516071"
        edge_width = (3.2 if is_article_context_edge else 1) + (float(attributes["confidence"]) * 2.4)
        edge_dashes = False if is_article_context_edge else True
        if compare_status == "new":
            edge_color = "#16a34a"
            edge_width += 1.4
            edge_dashes = False
        elif compare_status == "removed":
            edge_color = "#dc2626"
            edge_width = max(2.0, edge_width)
            edge_dashes = True
        network.add_edge(
            source_id,
            target_id,
            label=str(attributes["relationship"]) if show_visible_edge_labels else "",
            title="\n".join(tooltip_lines),
            value=2 + (float(attributes["confidence"]) * 4),
            width=edge_width,
            color=edge_color,
            dashes=edge_dashes,
            font={"size": 11 if show_visible_edge_labels else 8, "align": "middle"},
        )

    network.set_options(
        """
        const options = {
          "interaction": {
            "hover": true,
            "navigationButtons": true,
            "tooltipDelay": 120
          },
          "edges": {
            "font": {
              "size": 8
            },
            "smooth": {
              "enabled": true,
              "type": "dynamic"
            }
          },
          "layout": {
            "improvedLayout": true
          },
          "nodes": {
            "shadow": {
              "enabled": true,
              "color": "rgba(23, 32, 51, 0.12)",
              "size": 12,
              "x": 0,
              "y": 4
            }
          },
          "physics": {
            "barnesHut": {
              "gravitationalConstant": -5200,
              "springLength": 175
            },
            "minVelocity": 0.75
          }
        }
        """
    )

    network_html = network.generate_html()
    fit_script = """
    <script type="text/javascript">
	    setTimeout(function() {
	      if (typeof network !== "undefined") {
	        const fitGraph = function () {
	          network.fit({animation: false});
	        };
	        network.once("stabilizationIterationsDone", fitGraph);
	        setTimeout(fitGraph, 1200);
	      }
	    }, 0);
    </script>
    """
    network_html = network_html.replace("</body>", fit_script + "\n</body>")

    components.html(network_html, height=980, scrolling=True)


def inject_dashboard_css() -> None:
    """Apply a custom visual treatment to the Streamlit dashboard."""

    st.markdown(
        """
        <style>
        .stApp {
            background:
                radial-gradient(circle at top left, rgba(29, 78, 216, 0.10), transparent 28%),
                radial-gradient(circle at top right, rgba(15, 118, 110, 0.10), transparent 26%),
                linear-gradient(180deg, #f7f3eb 0%, #f9f7f2 36%, #f4efe5 100%);
            color: #172033;
        }
        .block-container {
            max-width: 1480px;
            padding-top: 2.1rem;
            padding-bottom: 3rem;
        }
        h1, h2, h3 {
            font-family: "Iowan Old Style", "Palatino Linotype", "Book Antiqua", Georgia, serif;
            letter-spacing: -0.02em;
            color: #172033;
        }
        h1 {
            font-size: 3rem;
            margin-bottom: 0.3rem;
        }
        .stMarkdown, [data-testid="stMetricLabel"], [data-testid="stMetricValue"] {
            font-family: "Avenir Next", "Segoe UI", "Trebuchet MS", sans-serif;
        }
        .stApp {
            color: #172033;
        }
        [data-testid="stSidebar"] {
            background: linear-gradient(180deg, rgba(255,255,255,0.94), rgba(244,239,229,0.95));
            border-right: 1px solid rgba(23, 32, 51, 0.08);
        }
        [data-testid="stSidebar"] [data-testid="stMarkdown"] h2,
        [data-testid="stSidebar"] [data-testid="stMarkdown"] h3 {
            color: #172033;
        }
        [data-baseweb="select"] > div,
        .stTextInput > div > div,
        .stTextArea textarea {
            background: #ffffff !important;
            border: 1px solid rgba(23, 32, 51, 0.18) !important;
            border-radius: 10px !important;
            color: #172033 !important;
            caret-color: #172033 !important;
            box-shadow: none !important;
        }
        .stTextInput input,
        .stTextArea textarea,
        .stTextArea textarea::placeholder,
        [data-baseweb="select"] span {
            color: #172033 !important;
            opacity: 1 !important;
        }
        .stTextArea label,
        .stCheckbox label,
        .stMarkdown p,
        [data-testid="stCaptionContainer"],
        [data-testid="stCaptionContainer"] p {
            color: #334155 !important;
        }
        .stButton button,
        [data-testid="stFormSubmitButton"] button {
            border-radius: 10px !important;
            border: 1px solid rgba(23, 32, 51, 0.16) !important;
            background: #ffffff !important;
            color: #172033 !important;
            box-shadow: none !important;
        }
        .stButton button p,
        [data-testid="stFormSubmitButton"] button p {
            color: inherit !important;
        }
        .stButton button:hover,
        [data-testid="stFormSubmitButton"] button:hover {
            border-color: rgba(29, 78, 216, 0.45) !important;
            color: #1d4ed8 !important;
        }
        [data-testid="stFormSubmitButton"] button[kind="primary"],
        [data-testid="stFormSubmitButton"] button {
            background: #1d4ed8 !important;
            border-color: #1d4ed8 !important;
            color: #ffffff !important;
        }
        [data-testid="stFormSubmitButton"] button p {
            color: #ffffff !important;
        }
        [data-testid="stChatMessage"] {
            background: #ffffff !important;
            border: 1px solid rgba(23, 32, 51, 0.10) !important;
            border-radius: 12px !important;
            box-shadow: 0 8px 24px rgba(23, 32, 51, 0.06) !important;
            margin: 0.75rem 0 !important;
            padding: 0.85rem 1rem !important;
        }
        [data-testid="stChatMessage"] * {
            color: #172033 !important;
        }
        [data-testid="stChatMessageAvatarUser"],
        [data-testid="stChatMessageAvatarAssistant"] {
            background: rgba(29, 78, 216, 0.10) !important;
            border-radius: 10px !important;
        }
        [data-testid="stExpander"] {
            background: #ffffff !important;
            border: 1px solid rgba(23, 32, 51, 0.10) !important;
            border-radius: 10px !important;
        }
        [data-testid="stDataFrame"] {
            color: #172033 !important;
        }
        [data-testid="stTabs"] {
            background: rgba(255, 255, 255, 0.65);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 12px;
            padding: 0.4rem 0.6rem 1rem 0.6rem;
            backdrop-filter: blur(8px);
        }
        [data-testid="stTabs"] button[role="tab"] {
            border-radius: 999px;
            padding: 0.7rem 1rem;
            color: #556277;
        }
        [data-testid="stTabs"] button[aria-selected="true"] {
            background: linear-gradient(135deg, #172033 0%, #1d4ed8 100%);
            color: #ffffff !important;
        }
        [data-testid="stTabs"] button[aria-selected="true"] p,
        [data-testid="stTabs"] button[aria-selected="true"] span {
            color: #ffffff !important;
        }
        .hero-panel {
            background:
                linear-gradient(135deg, rgba(23, 32, 51, 0.96), rgba(29, 78, 216, 0.90)),
                linear-gradient(135deg, #172033, #1d4ed8);
            color: #f8fafc;
            border-radius: 28px;
            padding: 1.6rem 1.8rem;
            box-shadow: 0 24px 60px rgba(23, 32, 51, 0.18);
            margin-bottom: 1rem;
        }
        .hero-kicker {
            text-transform: uppercase;
            letter-spacing: 0.18em;
            font-size: 0.74rem;
            opacity: 0.72;
            margin-bottom: 0.45rem;
        }
        .hero-title {
            font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
            font-size: 2.4rem;
            line-height: 1;
            margin-bottom: 0.4rem;
        }
        .hero-copy {
            max-width: 62rem;
            font-size: 1rem;
            color: rgba(248, 250, 252, 0.86);
            margin-bottom: 1rem;
        }
        .hero-badges, .pill-row {
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem;
        }
        .hero-badge, .signal-pill {
            display: inline-flex;
            align-items: center;
            border-radius: 999px;
            padding: 0.35rem 0.72rem;
            font-size: 0.82rem;
            font-weight: 600;
            border: 1px solid rgba(255, 255, 255, 0.14);
            background: rgba(255, 255, 255, 0.10);
            color: #f8fafc;
        }
        .overview-card, .panel-card, .list-card {
            background: rgba(255, 255, 255, 0.82);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 24px;
            padding: 1rem 1.1rem;
            box-shadow: 0 16px 40px rgba(23, 32, 51, 0.08);
        }
        .metric-card {
            background: linear-gradient(180deg, rgba(255,255,255,0.94), rgba(247,243,235,0.92));
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 22px;
            padding: 1rem 1.1rem;
            min-height: 8.5rem;
            box-shadow: 0 14px 32px rgba(23, 32, 51, 0.08);
        }
        .metric-label {
            color: #556277;
            text-transform: uppercase;
            letter-spacing: 0.12em;
            font-size: 0.74rem;
            margin-bottom: 0.45rem;
        }
        .metric-value {
            font-size: 2rem;
            line-height: 1;
            font-weight: 700;
            color: #172033;
            margin-bottom: 0.35rem;
        }
        .metric-note {
            color: #6a7485;
            font-size: 0.9rem;
        }
        .section-kicker {
            text-transform: uppercase;
            letter-spacing: 0.16em;
            color: #7c3f00;
            font-size: 0.72rem;
            margin-bottom: 0.2rem;
        }
        .section-title {
            font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
            font-size: 1.6rem;
            color: #172033;
            margin-bottom: 0.2rem;
        }
        .section-copy {
            color: #5d6676;
            margin-bottom: 0.6rem;
        }
        .meta-grid {
            display: grid;
            grid-template-columns: repeat(2, minmax(0, 1fr));
            gap: 0.7rem;
            margin-bottom: 0.75rem;
        }
        .meta-card {
            background: rgba(247, 243, 235, 0.88);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 18px;
            padding: 0.85rem 0.95rem;
        }
        .meta-label {
            color: #6a7485;
            font-size: 0.72rem;
            text-transform: uppercase;
            letter-spacing: 0.12em;
            margin-bottom: 0.25rem;
        }
        .meta-value {
            color: #172033;
            font-size: 1rem;
            font-weight: 600;
        }
        .summary-card {
            background: linear-gradient(135deg, rgba(29, 78, 216, 0.08), rgba(15, 118, 110, 0.08));
            border: 1px solid rgba(29, 78, 216, 0.10);
            border-radius: 22px;
            padding: 1rem 1.1rem;
            margin-bottom: 0.9rem;
        }
        .summary-title {
            color: #1d4ed8;
            font-size: 0.74rem;
            text-transform: uppercase;
            letter-spacing: 0.14em;
            margin-bottom: 0.35rem;
        }
        .summary-body {
            color: #172033;
            font-size: 1rem;
            line-height: 1.6;
        }
        .signal-group {
            margin-top: 0.85rem;
            padding-top: 0.85rem;
            border-top: 1px solid rgba(23, 32, 51, 0.08);
        }
        .signal-label {
            color: #5d6676;
            text-transform: uppercase;
            letter-spacing: 0.12em;
            font-size: 0.72rem;
            margin-bottom: 0.45rem;
        }
        .signal-pill {
            background: rgba(23, 32, 51, 0.06);
            color: #172033;
            border: 1px solid rgba(23, 32, 51, 0.08);
        }
        .article-source-divider {
            height: 1px;
            background: rgba(23, 32, 51, 0.08);
            margin: 1.25rem 0 0.85rem 0;
        }
        .media-shell {
            background: linear-gradient(180deg, rgba(255,255,255,0.88), rgba(244,239,229,0.84));
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 22px;
            padding: 0.9rem 1rem 0.55rem 1rem;
            box-shadow: 0 14px 34px rgba(23, 32, 51, 0.07);
            margin: 0.05rem 0 0.85rem 0;
        }
        .media-intro {
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            justify-content: space-between;
            gap: 0.7rem;
            margin-bottom: 0.65rem;
        }
        .media-title {
            font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
            font-size: 1.24rem;
            color: #172033;
        }
        .media-note {
            color: #5d6676;
            font-size: 0.92rem;
            max-width: 34rem;
        }
        .media-chip-row {
            display: flex;
            flex-wrap: wrap;
            gap: 0.45rem;
            margin-bottom: 0.35rem;
        }
        .media-chip {
            display: inline-flex;
            align-items: center;
            gap: 0.35rem;
            border-radius: 999px;
            padding: 0.38rem 0.72rem;
            background: rgba(23, 32, 51, 0.06);
            border: 1px solid rgba(23, 32, 51, 0.08);
            color: #172033;
            font-size: 0.82rem;
            font-weight: 600;
        }
        .image-rail {
            background: rgba(248, 250, 252, 0.82);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 22px;
            padding: 0.85rem;
            margin-bottom: 0.95rem;
        }
        .image-rail-title {
            color: #7c3f00;
            text-transform: uppercase;
            letter-spacing: 0.14em;
            font-size: 0.72rem;
            margin-bottom: 0.35rem;
        }
        .image-rail-copy {
            color: #5d6676;
            font-size: 0.92rem;
            margin-bottom: 0.75rem;
        }
        .image-empty {
            border: 1px dashed rgba(23, 32, 51, 0.18);
            border-radius: 18px;
            padding: 1rem;
            color: #5d6676;
            background: rgba(255,255,255,0.72);
        }
        .legend-strip {
            display: flex;
            flex-wrap: wrap;
            gap: 0.6rem;
            margin: 0.4rem 0 0.8rem 0;
        }
        .legend-chip {
            display: inline-flex;
            align-items: center;
            gap: 0.45rem;
            border-radius: 999px;
            padding: 0.42rem 0.78rem;
            background: rgba(255,255,255,0.82);
            border: 1px solid rgba(23, 32, 51, 0.08);
            font-size: 0.84rem;
        }
        .graph-shell {
            background: rgba(255, 255, 255, 0.84);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 26px;
            padding: 1.05rem 1.1rem 1.15rem 1.1rem;
            box-shadow: 0 18px 42px rgba(23, 32, 51, 0.08);
            margin-top: 0.95rem;
        }
        .graph-stage-title {
            font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
            font-size: 1.4rem;
            color: #172033;
            margin-bottom: 0.2rem;
        }
        .graph-stage-copy {
            color: #5d6676;
            margin-bottom: 0.7rem;
        }
        .graph-control-card {
            background: rgba(255, 255, 255, 0.86);
            border: 1px solid rgba(23, 32, 51, 0.08);
            border-radius: 24px;
            padding: 1rem 1rem 0.55rem 1rem;
            box-shadow: 0 14px 32px rgba(23, 32, 51, 0.07);
            margin-top: 0.6rem;
            margin-bottom: 0.75rem;
        }
        .graph-control-title {
            font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
            font-size: 1.28rem;
            color: #172033;
            margin-bottom: 0.2rem;
        }
        .graph-control-copy {
            color: #5d6676;
            margin-bottom: 0.8rem;
        }
        .legend-dot {
            width: 0.7rem;
            height: 0.7rem;
            border-radius: 999px;
            display: inline-block;
        }
        /* ── Status badges ─────────────────────────────────────────── */
        .status-badge {
            display: inline-flex;
            align-items: center;
            border-radius: 999px;
            padding: 0.28rem 0.68rem;
            font-size: 0.76rem;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            white-space: nowrap;
        }
        .status-emerging   { background: rgba(15,118,110,0.12);  color: #0f766e; border: 1px solid rgba(15,118,110,0.22); }
        .status-strengthening { background: rgba(29,78,216,0.12); color: #1d4ed8; border: 1px solid rgba(29,78,216,0.22); }
        .status-weakening  { background: rgba(217,119,6,0.12);   color: #b45309; border: 1px solid rgba(217,119,6,0.22); }
        .status-stable     { background: rgba(107,114,128,0.12); color: #4b5563; border: 1px solid rgba(107,114,128,0.22); }
        .status-reversed   { background: rgba(220,38,38,0.12);   color: #dc2626; border: 1px solid rgba(220,38,38,0.22); }
        .status-transitioning { background: rgba(124,63,0,0.12); color: #7c3f00; border: 1px solid rgba(124,63,0,0.22); }
        /* Trend status */
        .trend-new          { background: rgba(15,118,110,0.12);  color: #0f766e; border: 1px solid rgba(15,118,110,0.22); }
        .trend-strengthening{ background: rgba(29,78,216,0.12);   color: #1d4ed8; border: 1px solid rgba(29,78,216,0.22); }
        .trend-recurring    { background: rgba(107,114,128,0.12); color: #4b5563; border: 1px solid rgba(107,114,128,0.22); }
        .trend-weakening    { background: rgba(217,119,6,0.12);   color: #b45309; border: 1px solid rgba(217,119,6,0.22); }
        /* Entity type badges */
        .etype-badge {
            display: inline-block;
            border-radius: 8px;
            padding: 0.18rem 0.52rem;
            font-size: 0.75rem;
            font-weight: 600;
            text-transform: capitalize;
        }
        .etype-person       { background: rgba(29,78,216,0.10);  color: #1d4ed8; border: 1px solid rgba(29,78,216,0.18); }
        .etype-company      { background: rgba(15,118,110,0.10); color: #0f766e; border: 1px solid rgba(15,118,110,0.18); }
        .etype-organization { background: rgba(15,118,110,0.10); color: #0f766e; border: 1px solid rgba(15,118,110,0.18); }
        .etype-country      { background: rgba(124,63,0,0.10);   color: #7c3f00; border: 1px solid rgba(124,63,0,0.18); }
        .etype-technology   { background: rgba(109,40,217,0.10); color: #6d28d9; border: 1px solid rgba(109,40,217,0.18); }
        .etype-default      { background: rgba(23,32,51,0.08);   color: #556277; border: 1px solid rgba(23,32,51,0.14); }
        /* Narrative card */
        .narrative-card {
            background: rgba(255,255,255,0.88);
            border: 1px solid rgba(23,32,51,0.08);
            border-radius: 20px;
            padding: 1rem 1.15rem;
            margin-bottom: 0.65rem;
            box-shadow: 0 8px 24px rgba(23,32,51,0.06);
            transition: box-shadow 0.18s;
        }
        .narrative-card:hover {
            box-shadow: 0 14px 36px rgba(23,32,51,0.11);
        }
        .narrative-card-header {
            display: flex;
            align-items: flex-start;
            justify-content: space-between;
            gap: 0.7rem;
            margin-bottom: 0.45rem;
        }
        .narrative-card-title {
            font-family: "Iowan Old Style","Palatino Linotype",Georgia,serif;
            font-size: 1.04rem;
            font-weight: 700;
            color: #172033;
            flex: 1;
            line-height: 1.35;
        }
        .narrative-thesis {
            color: #5d6676;
            font-size: 0.91rem;
            line-height: 1.57;
            margin-bottom: 0.55rem;
        }
        .narrative-footer {
            display: flex;
            flex-wrap: wrap;
            gap: 0.35rem 0.6rem;
            align-items: center;
            font-size: 0.79rem;
            color: #6a7485;
            margin-top: 0.4rem;
        }
        /* Importance score bar */
        .score-bar-outer {
            background: rgba(23,32,51,0.08);
            border-radius: 999px;
            height: 5px;
            width: 100%;
            margin: 0.4rem 0 0.35rem 0;
        }
        .score-bar-inner {
            background: linear-gradient(90deg, #1d4ed8 0%, #0f766e 100%);
            border-radius: 999px;
            height: 5px;
        }
        /* Sidebar db stats card */
        .sidebar-db-card {
            background: rgba(255,255,255,0.72);
            border: 1px solid rgba(23,32,51,0.08);
            border-radius: 16px;
            padding: 0.75rem 0.9rem;
            margin-top: 0.3rem;
        }
        .sidebar-db-kicker {
            font-size: 0.7rem;
            text-transform: uppercase;
            letter-spacing: 0.14em;
            color: #7c3f00;
            margin-bottom: 0.5rem;
        }
        .sidebar-stat-row {
            display: flex;
            align-items: center;
            justify-content: space-between;
            padding: 0.18rem 0;
            font-size: 0.84rem;
        }
        .sidebar-stat-label { color: #5d6676; }
        .sidebar-stat-value { font-weight: 700; color: #172033; }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        """
        <style media="not all" data-disabled="command-center">
        :root {
            --na-ink: #0f172a;
            --na-muted: #64748b;
            --na-panel: #ffffff;
            --na-line: #dbe3ef;
            --na-blue: #2563eb;
            --na-teal: #0f766e;
            --na-amber: #b45309;
            --na-red: #dc2626;
            --na-violet: #7c3aed;
            --na-bg: #f4f7fb;
            --na-rail: #0f172a;
        }
        .stApp {
            background: var(--na-bg) !important;
            color: var(--na-ink) !important;
            font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
        }
        .block-container {
            max-width: 1320px;
            padding-top: 1.5rem;
            padding-bottom: 3rem;
        }
        h1, h2, h3,
        .hero-title,
        .section-title,
        .graph-stage-title,
        .graph-control-title,
        .media-title,
        .narrative-card-title {
            font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif !important;
            letter-spacing: 0 !important;
            color: var(--na-ink);
        }
        [data-testid="stSidebar"] {
            background: var(--na-rail) !important;
            border-right: 1px solid #1e293b !important;
        }
        [data-testid="stSidebar"] * {
            color: #e2e8f0;
        }
        [data-testid="stSidebar"] [data-testid="stCaptionContainer"],
        [data-testid="stSidebar"] [data-testid="stCaptionContainer"] p,
        [data-testid="stSidebar"] .stMarkdown p {
            color: #94a3b8 !important;
        }
        .sidebar-brand {
            margin-bottom: 0.9rem;
        }
        .sidebar-brand-title {
            color: #f8fafc;
            font-size: 1.55rem;
            line-height: 1.25;
            font-weight: 800;
        }
        .sidebar-brand-copy {
            color: #94a3b8;
            font-size: 0.78rem;
            line-height: 1.35;
        }
        .sidebar-source-card {
            display: flex;
            align-items: center;
            gap: 0.65rem;
            background: #12213a;
            border: 1px solid #243b5a;
            border-radius: 8px;
            padding: 0.8rem;
            margin: 0.5rem 0 1rem;
        }
        .sidebar-source-dot {
            width: 0.62rem;
            height: 0.62rem;
            flex: 0 0 auto;
            border-radius: 999px;
            background: #2dd4bf;
        }
        .sidebar-source-label {
            color: #94a3b8;
            font-size: 0.74rem;
            line-height: 1.1;
        }
        .sidebar-source-value {
            color: #f8fafc;
            font-size: 0.95rem;
            font-weight: 800;
            line-height: 1.2;
        }
        [data-testid="stSidebar"] [data-baseweb="select"] > div,
        [data-testid="stSidebar"] .stTextInput > div > div {
            background: #172a45 !important;
            border: 1px solid #254465 !important;
            border-radius: 7px !important;
            color: #e2e8f0 !important;
        }
        [data-testid="stSidebar"] input,
        [data-testid="stSidebar"] [data-baseweb="select"] span {
            color: #e2e8f0 !important;
        }
        [data-testid="stSidebar"] label,
        [data-testid="stSidebar"] label p {
            color: #7dd3fc !important;
            font-size: 0.72rem !important;
            font-weight: 800 !important;
            text-transform: uppercase;
            letter-spacing: 0.08em;
        }
        .sidebar-db-card {
            background: #f8fafc !important;
            border: 0 !important;
            border-radius: 8px !important;
            padding: 0.9rem !important;
        }
        .sidebar-db-kicker {
            color: #475569 !important;
            letter-spacing: 0;
            text-transform: none;
            font-size: 0.94rem;
            font-weight: 800;
        }
        .sidebar-stat-row {
            font-size: 0.8rem;
            padding: 0.23rem 0;
        }
        .sidebar-stat-label { color: #0f172a !important; }
        .sidebar-stat-value { color: #0f172a !important; }
        .na-command-bar {
            display: grid;
            grid-template-columns: minmax(18rem, 1fr) minmax(18rem, 23rem) auto auto;
            gap: 0.9rem;
            align-items: center;
            min-height: 5.4rem;
            background: #ffffff;
            border: 1px solid var(--na-line);
            border-radius: 8px;
            box-shadow: 0 14px 32px rgba(15, 23, 42, 0.06);
            padding: 1.1rem 1.35rem;
            margin-bottom: 1.25rem;
        }
        .na-command-title {
            font-size: 1.32rem;
            line-height: 1.28;
            font-weight: 850;
            color: var(--na-ink);
        }
        .na-command-subtitle {
            color: #7b8da8;
            font-size: 0.78rem;
            line-height: 1.35;
            margin-top: 0.15rem;
        }
        .na-command-query {
            min-height: 2.85rem;
            display: flex;
            align-items: center;
            border: 1px solid #cbd5e1;
            background: #f8fafc;
            border-radius: 8px;
            padding: 0 0.9rem;
            color: var(--na-muted);
            font-size: 0.9rem;
            overflow: hidden;
            white-space: nowrap;
            text-overflow: ellipsis;
        }
        .na-command-button {
            display: inline-flex;
            align-items: center;
            justify-content: center;
            min-height: 2.5rem;
            border-radius: 8px;
            padding: 0 1rem;
            color: #ffffff;
            background: #111827;
            font-size: 0.82rem;
            font-weight: 850;
            white-space: nowrap;
        }
        .na-command-button.teal {
            background: var(--na-teal);
        }
        .metric-card {
            background: #ffffff !important;
            border: 1px solid var(--na-line) !important;
            border-radius: 8px !important;
            box-shadow: 0 14px 32px rgba(15, 23, 42, 0.05) !important;
            min-height: 7.25rem !important;
            padding: 1rem !important;
        }
        .metric-label {
            display: flex;
            align-items: center;
            justify-content: space-between;
            color: var(--na-muted) !important;
            letter-spacing: 0.06em !important;
            font-size: 0.68rem !important;
            font-weight: 850 !important;
            margin-bottom: 0.55rem;
        }
        .metric-dot {
            width: 0.55rem;
            height: 0.55rem;
            border-radius: 999px;
            background: var(--metric-accent, var(--na-blue));
        }
        .metric-value {
            color: var(--na-ink) !important;
            font-size: 2rem !important;
            font-weight: 850 !important;
            line-height: 1 !important;
        }
        .metric-note {
            color: #475569 !important;
            font-size: 0.78rem !important;
            margin-top: 0.45rem;
        }
        .na-command-center {
            display: grid;
            grid-template-columns: minmax(32rem, 1.55fr) minmax(19rem, 0.9fr);
            gap: 1.5rem;
            margin: 1.35rem 0 1.5rem;
        }
        .na-panel {
            background: #ffffff;
            border: 1px solid var(--na-line);
            border-radius: 8px;
            box-shadow: 0 14px 32px rgba(15, 23, 42, 0.06);
            padding: 1.5rem;
            min-width: 0;
        }
        .na-panel.dark {
            background: #111827;
            border-color: #243244;
            color: #f8fafc;
        }
        .na-panel-head {
            display: flex;
            align-items: flex-start;
            justify-content: space-between;
            gap: 1rem;
            margin-bottom: 1.2rem;
        }
        .na-panel-title {
            color: inherit;
            font-size: 1.12rem;
            line-height: 1.25;
            font-weight: 850;
        }
        .na-panel-copy {
            color: #7b8da8;
            font-size: 0.78rem;
            line-height: 1.35;
            margin-top: 0.15rem;
        }
        .na-panel.dark .na-panel-copy {
            color: #94a3b8;
        }
        .na-chip-row {
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem;
            align-items: center;
        }
        .na-chip {
            display: inline-flex;
            align-items: center;
            min-height: 1.75rem;
            border-radius: 7px;
            padding: 0 0.65rem;
            background: #eff6ff;
            color: #1d4ed8;
            font-size: 0.7rem;
            font-weight: 850;
            white-space: nowrap;
        }
        .na-chip.teal { background: #ecfdf5; color: #047857; }
        .na-chip.amber { background: #fef3c7; color: #92400e; }
        .na-chip.violet { background: #f5f3ff; color: #6d28d9; }
        .na-chip.slate { background: #f1f5f9; color: #475569; }
        .na-graph-stage {
            height: 24rem;
            border-radius: 8px;
            border: 1px solid var(--na-line);
            background: #f8fafc;
            overflow: hidden;
        }
        .na-graph-stage svg {
            width: 100%;
            height: 100%;
            display: block;
        }
        .na-question-card {
            border: 1px solid #334155;
            background: #1e293b;
            border-radius: 8px;
            padding: 1rem;
            margin-bottom: 1.15rem;
        }
        .na-question-label,
        .na-answer-label {
            display: block;
            color: #7dd3fc;
            font-size: 0.64rem;
            font-weight: 900;
            letter-spacing: 0.06em;
            text-transform: uppercase;
            margin-bottom: 0.5rem;
        }
        .na-question-card strong {
            color: #f8fafc;
            font-size: 1rem;
            line-height: 1.25;
        }
        .na-answer-card {
            background: #f8fafc;
            color: var(--na-ink);
            border-radius: 8px;
            padding: 1rem;
            margin-bottom: 1.15rem;
        }
        .na-answer-label {
            color: #475569;
        }
        .na-answer-card p {
            color: var(--na-ink) !important;
            font-size: 0.95rem;
            line-height: 1.34;
            margin: 0;
        }
        .na-lower-grid {
            display: grid;
            grid-template-columns: repeat(3, minmax(0, 1fr));
            gap: 1.5rem;
            margin-bottom: 1.5rem;
        }
        .na-trend-row {
            display: grid;
            grid-template-columns: 6.8rem 1fr 2.2rem;
            gap: 0.65rem;
            align-items: center;
            min-height: 1.9rem;
            font-size: 0.78rem;
        }
        .na-trend-bar {
            height: 0.5rem;
            border-radius: 999px;
            background: #e2e8f0;
            overflow: hidden;
        }
        .na-trend-fill {
            height: 100%;
            border-radius: inherit;
            background: var(--bar-color, var(--na-teal));
            width: var(--bar-width, 0%);
        }
        .na-article-title {
            margin-top: 1.05rem;
            color: var(--na-ink);
            font-size: 1.25rem;
            line-height: 1.25;
            font-weight: 850;
        }
        .na-article-copy {
            color: #334155 !important;
            font-size: 0.84rem;
            line-height: 1.45;
            margin: 0.65rem 0 1rem;
        }
        .na-pipeline-row {
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 0.75rem;
            border-top: 1px solid #eef2f7;
            padding-top: 0.65rem;
            margin-top: 0.65rem;
            font-size: 0.8rem;
        }
        [data-testid="stTabs"] {
            background: transparent !important;
            border: 0 !important;
            border-radius: 0 !important;
            padding: 0 !important;
            backdrop-filter: none !important;
        }
        [data-testid="stTabs"] button[role="tab"] {
            border-radius: 8px !important;
            padding: 0.65rem 0.85rem !important;
            color: #475569 !important;
            font-weight: 750;
        }
        [data-testid="stTabs"] button[aria-selected="true"] {
            background: #111827 !important;
            color: #ffffff !important;
        }
        .hero-panel {
            background: #ffffff !important;
            color: var(--na-ink) !important;
            border: 1px solid var(--na-line) !important;
            border-radius: 8px !important;
            box-shadow: 0 14px 32px rgba(15, 23, 42, 0.06) !important;
            padding: 1.15rem 1.3rem !important;
        }
        .hero-kicker,
        .section-kicker {
            color: #475569 !important;
            letter-spacing: 0.06em !important;
        }
        .hero-title {
            font-size: 1.55rem !important;
            color: var(--na-ink) !important;
        }
        .hero-copy,
        .section-copy {
            color: #64748b !important;
        }
        .hero-badge,
        .signal-pill {
            border-radius: 7px !important;
            border: 0 !important;
            background: #eff6ff !important;
            color: #1d4ed8 !important;
            font-size: 0.76rem !important;
            font-weight: 800 !important;
        }
        .overview-card, .panel-card, .list-card, .summary-card, .meta-card,
        .graph-shell, .graph-control-card, .narrative-card, .media-shell, .image-rail {
            border-radius: 8px !important;
            background: #ffffff !important;
            border: 1px solid var(--na-line) !important;
            box-shadow: 0 14px 32px rgba(15, 23, 42, 0.05) !important;
        }
        .stButton button,
        [data-testid="stFormSubmitButton"] button {
            border-radius: 8px !important;
            font-weight: 800 !important;
        }
        [data-testid="stFormSubmitButton"] button {
            background: var(--na-teal) !important;
            border-color: var(--na-teal) !important;
        }
        [data-testid="stChatMessage"] {
            border-radius: 8px !important;
            box-shadow: none !important;
        }
        @media (max-width: 1100px) {
            .na-command-bar,
            .na-command-center,
            .na-lower-grid {
                grid-template-columns: 1fr;
            }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def _render_metric_card(label: str, value: str, note: str) -> None:
    st.markdown(
        f"""
        <div class="metric-card">
          <div class="metric-label">{escape(label)}</div>
          <div class="metric-value">{escape(value)}</div>
          <div class="metric-note">{escape(note)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_section_intro(title: str, copy: str, kicker: str = "Briefing") -> None:
    st.markdown(
        f"""
        <div class="section-kicker">{escape(kicker)}</div>
        <div class="section-title">{escape(title)}</div>
        <div class="section-copy">{escape(copy)}</div>
        """,
        unsafe_allow_html=True,
    )


def _render_command_bar(selected_week: str, total_filtered: int, avg_importance: float, search_term: str) -> None:
    scope_label = selected_week if selected_week != "All" else "All indexed weeks"
    query_label = search_term or "Ask about entities, relationships, or weekly shifts"
    st.markdown(
        f"""
        <div class="na-command-bar">
          <div>
            <div class="na-command-title">Narrative Intelligence Console</div>
            <div class="na-command-subtitle">
              Latest scope: {escape(scope_label)} | {total_filtered} article(s) | average importance {avg_importance:.1f}
            </div>
          </div>
          <div class="na-command-query">{escape(query_label)}</div>
          <div class="na-command-button teal">Deep Search</div>
          <div class="na-command-button">Run Pipeline</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _first_nonempty_value(dataframe: pd.DataFrame, column: str, fallback: str) -> str:
    if dataframe.empty or column not in dataframe.columns:
        return fallback
    values = dataframe[column].dropna().astype(str)
    for value in values:
        stripped = value.strip()
        if stripped:
            return stripped
    return fallback


def _graph_preview_stats(
    preview_graph: dict[str, Any] | None,
    metrics: dict[str, int],
) -> tuple[int, int, str]:
    if preview_graph:
        return (
            len(preview_graph.get("nodes", [])),
            len(preview_graph.get("edges", [])),
            str(preview_graph.get("week") or "latest"),
        )
    return (
        int(metrics.get("entities", 0)) + int(metrics.get("themes", 0)) + int(metrics.get("narratives", 0)),
        int(metrics.get("graph_edges", 0)),
        "latest",
    )


def _render_graph_preview_svg(labels: dict[str, str]) -> str:
    label_openai = escape(labels["source"])
    label_narrative = escape(labels["narrative"])
    label_entity = escape(labels["entity"])
    label_theme = escape(labels["theme"])
    label_article = escape(labels["article"])
    return f"""
    <svg viewBox="0 0 700 384" role="img" aria-label="Narrative graph preview">
      <rect width="700" height="384" fill="#f8fafc"></rect>
      <g stroke="#cbd5e1" stroke-opacity=".35">
        <path d="M40 48 L170 318"></path>
        <path d="M120 48 L220 318"></path>
        <path d="M200 48 L290 318"></path>
        <path d="M280 48 L355 318"></path>
        <path d="M360 48 L430 318"></path>
        <path d="M440 48 L505 318"></path>
        <path d="M520 48 L580 318"></path>
      </g>
      <g fill="none" stroke-linecap="round" stroke-width="2.5">
        <path d="M126 130 L292 86" stroke="#2563eb" stroke-opacity=".65"></path>
        <path d="M292 86 L505 112" stroke="#2563eb" stroke-opacity=".45"></path>
        <path d="M126 130 L220 258" stroke="#0f766e" stroke-opacity=".65"></path>
        <path d="M220 258 L405 286" stroke="#0f766e" stroke-opacity=".45"></path>
        <path d="M505 112 L574 250" stroke="#dc2626" stroke-opacity=".55"></path>
        <path d="M292 86 L346 214" stroke="#7c3aed" stroke-opacity=".55"></path>
        <path d="M405 286 L574 250" stroke="#b45309" stroke-opacity=".5"></path>
      </g>
      <g font-family="Inter, system-ui, sans-serif" font-size="11" font-weight="750" text-anchor="middle" fill="#334155">
        <circle cx="116" cy="122" r="36" fill="#2563eb" opacity=".13"></circle>
        <circle cx="116" cy="122" r="28" fill="#2563eb" stroke="#fff" stroke-width="2"></circle>
        <text x="116" y="173">{label_openai}</text>
        <circle cx="292" cy="82" r="40" fill="#7c3aed" opacity=".13"></circle>
        <circle cx="292" cy="82" r="31" fill="#7c3aed" stroke="#fff" stroke-width="2"></circle>
        <text x="292" y="135">{label_narrative}</text>
        <circle cx="505" cy="105" r="31" fill="#0f766e" opacity=".13"></circle>
        <circle cx="505" cy="105" r="23" fill="#0f766e" stroke="#fff" stroke-width="2"></circle>
        <text x="505" y="150">{label_entity}</text>
        <circle cx="220" cy="256" r="33" fill="#b45309" opacity=".13"></circle>
        <circle cx="220" cy="256" r="25" fill="#b45309" stroke="#fff" stroke-width="2"></circle>
        <text x="220" y="304">{label_theme}</text>
        <circle cx="405" cy="286" r="35" fill="#dc2626" opacity=".13"></circle>
        <circle cx="405" cy="286" r="27" fill="#dc2626" stroke="#fff" stroke-width="2"></circle>
        <text x="405" y="338">Capital Wars</text>
        <circle cx="574" cy="250" r="39" fill="#f59e0b" opacity=".15"></circle>
        <circle cx="574" cy="250" r="30" fill="#f59e0b" stroke="#fff" stroke-width="2"></circle>
        <text x="574" y="306">Oil Supply Shock</text>
        <circle cx="574" cy="68" r="28" fill="#0f766e" opacity=".13"></circle>
        <circle cx="574" cy="68" r="21" fill="#0f766e" stroke="#fff" stroke-width="2"></circle>
        <text x="574" y="111">{label_article}</text>
      </g>
    </svg>
    """


def _render_command_center(
    metrics: dict[str, int],
    preview_graph: dict[str, Any] | None,
    filtered_articles: pd.DataFrame,
    filtered_narratives: pd.DataFrame,
    selected_week: str,
    search_term: str,
) -> None:
    node_count, edge_count, graph_week = _graph_preview_stats(preview_graph, metrics)
    article_label = _first_nonempty_value(filtered_articles, "title", "Alpha trial")
    narrative_label = _first_nonempty_value(filtered_narratives, "name", "AI Demand vs Supply")
    question = search_term or f"What changed in {selected_week if selected_week != 'All' else 'the latest graph'}?"
    answer = _first_sentence(
        _first_nonempty_value(
            filtered_articles,
            "summary",
            "Graph facts, article summaries, and citations are ready for the current scope.",
        ),
        185,
    )
    labels = {
        "source": "OpenAI",
        "narrative": narrative_label[:24],
        "entity": "Microsoft",
        "theme": "Fed",
        "article": article_label[:18],
    }
    graph_svg = _render_graph_preview_svg(labels)
    st.markdown(
        f"""
        <div class="na-command-center">
          <section class="na-panel">
            <div class="na-panel-head">
              <div>
                <div class="na-panel-title">Relationship Graph</div>
                <div class="na-panel-copy">Week-scoped evidence network | {escape(graph_week)}</div>
              </div>
              <div class="na-chip-row">
                <span class="na-chip">{node_count} nodes</span>
                <span class="na-chip teal">{edge_count} links</span>
                <span class="na-chip amber">Compare off</span>
              </div>
            </div>
            <div class="na-graph-stage">{graph_svg}</div>
          </section>
          <section class="na-panel dark">
            <div class="na-panel-head">
              <div>
                <div class="na-panel-title">Ask Graph</div>
                <div class="na-panel-copy">GraphRAG answer with citations</div>
              </div>
            </div>
            <div class="na-question-card">
              <span class="na-question-label">Question</span>
              <strong>{escape(question)}</strong>
            </div>
            <div class="na-answer-card">
              <span class="na-answer-label">Fast answer</span>
              <p>{escape(answer)}</p>
            </div>
            <div class="na-chip-row">
              <span class="na-chip">graph facts</span>
              <span class="na-chip teal">citations</span>
              <span class="na-chip amber">LLM optional</span>
            </div>
          </section>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _trend_status_counts(trend_df: pd.DataFrame, filtered_narratives: pd.DataFrame) -> dict[str, int]:
    if not trend_df.empty and "trend_status" in trend_df.columns:
        counts = trend_df["trend_status"].fillna("").astype(str).str.upper().value_counts().to_dict()
        return {
            "New": int(counts.get("NEW", 0)),
            "Strengthening": int(counts.get("STRENGTHENING", 0)),
            "Recurring": int(counts.get("RECURRING", 0)),
            "Weakening": int(counts.get("WEAKENING", 0)),
        }
    if not filtered_narratives.empty and "status" in filtered_narratives.columns:
        counts = filtered_narratives["status"].fillna("").astype(str).str.lower().value_counts().to_dict()
        return {
            "New": int(counts.get("emerging", 0)),
            "Strengthening": int(counts.get("strengthening", 0)),
            "Recurring": int(counts.get("stable", 0)),
            "Weakening": int(counts.get("weakening", 0)),
        }
    return {"New": 0, "Strengthening": 0, "Recurring": 0, "Weakening": 0}


def _render_trend_rows(counts: dict[str, int]) -> str:
    colors = {
        "New": "#0f766e",
        "Strengthening": "#2563eb",
        "Recurring": "#64748b",
        "Weakening": "#b45309",
    }
    max_count = max([*counts.values(), 1])
    rows = []
    for label in ["New", "Strengthening", "Recurring", "Weakening"]:
        value = counts.get(label, 0)
        width = max(4, int((value / max_count) * 100)) if value else 4
        rows.append(
            f'<div class="na-trend-row">'
            f'<span>{escape(label)}</span>'
            f'<div class="na-trend-bar">'
            f'<div class="na-trend-fill" style="--bar-width:{width}%;--bar-color:{colors[label]};"></div>'
            f'</div>'
            f'<strong>{value}</strong>'
            f'</div>'
        )
    return "".join(rows)


def _render_lower_command_panels(
    metrics: dict[str, int],
    filtered_articles: pd.DataFrame,
    filtered_narratives: pd.DataFrame,
    trend_df: pd.DataFrame,
    selected_week: str,
    data_backend: str,
) -> None:
    trend_counts = _trend_status_counts(trend_df, filtered_narratives)
    if filtered_articles.empty:
        article_title = "No article selected"
        article_summary = "Adjust filters to inspect an article extraction."
        article_category = "None"
        article_importance = "0"
    else:
        article = filtered_articles.sort_values(
            by=["importance_score", "published_date"],
            ascending=[False, False],
        ).iloc[0]
        article_title = str(article.get("title") or "Untitled")
        article_summary = _first_sentence(str(article.get("summary") or ""), 170)
        article_category = str(article.get("category") or "Uncategorized")
        article_importance = str(article.get("importance_score") or "0")
    trend_rows = _render_trend_rows(trend_counts)
    st.markdown(
        f"""
        <div class="na-lower-grid">
          <section class="na-panel">
            <div class="na-panel-title">Narrative Movement</div>
            <div class="na-panel-copy">{escape(selected_week)} trend status</div>
            <div style="height:1rem;"></div>
            {trend_rows}
          </section>
          <section class="na-panel">
            <div class="na-panel-title">Article Workbench</div>
            <div class="na-panel-copy">Selected extraction</div>
            <div class="na-article-title">{escape(article_title)}</div>
            <p class="na-article-copy">{escape(article_summary)}</p>
            <div class="na-chip-row">
              <span class="na-chip">{escape(article_category)}</span>
              <span class="na-chip amber">Importance {escape(article_importance)}</span>
              <span class="na-chip violet">{len(filtered_narratives)} narratives</span>
            </div>
          </section>
          <section class="na-panel">
            <div class="na-panel-title">Pipeline Health</div>
            <div class="na-panel-copy">Current dashboard readout</div>
            <div class="na-pipeline-row"><span>Graph source</span><span class="na-chip teal">{escape(data_backend)}</span></div>
            <div class="na-pipeline-row"><span>Articles in scope</span><span class="na-chip">{len(filtered_articles)}</span></div>
            <div class="na-pipeline-row"><span>Total graph edges</span><span class="na-chip">{metrics.get('graph_edges', 0)}</span></div>
            <div class="na-pipeline-row"><span>Registry</span><span class="na-chip teal">Ready</span></div>
          </section>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_signal_group(label: str, items: list[str], empty_label: str) -> None:
    pills = "".join(f'<span class="signal-pill">{escape(item)}</span>' for item in items) or (
        f'<span class="signal-pill">{escape(empty_label)}</span>'
    )
    st.markdown(
        f"""
        <div class="signal-group">
          <div class="signal-label">{escape(label)}</div>
          <div class="pill-row">{pills}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _status_badge_html(status: str, badge_classes: dict[str, str]) -> str:
    """Return an HTML badge span for a status value."""
    css_class = badge_classes.get(str(status).lower(), "status-stable")
    return f'<span class="status-badge {css_class}">{escape(str(status))}</span>'


def _entity_type_badge_html(entity_type: str) -> str:
    """Return an HTML badge span for an entity type."""
    css_class = ENTITY_TYPE_CLASSES.get(str(entity_type).lower(), "etype-default")
    return f'<span class="etype-badge {css_class}">{escape(str(entity_type))}</span>'


def _render_narrative_card(row: pd.Series) -> None:
    """Render a single narrative as a styled card."""
    name = str(row.get("name", "") or "")
    thesis = str(row.get("thesis", "") or "No thesis recorded.")
    status = str(row.get("status", "") or "")
    importance = float(row.get("importance_score", 0) or 0)
    mentions = int(row.get("mention_count", 0) or 0)
    first_seen = str(row.get("first_seen_date", "") or "")
    last_seen = str(row.get("last_seen_date", "") or "")
    badge_html = _status_badge_html(status, STATUS_BADGE_CLASSES)
    score_pct = min(100, int(importance * 10))
    date_parts = []
    if first_seen:
        date_parts.append(f"First: {escape(first_seen)}")
    if last_seen:
        date_parts.append(f"Last: {escape(last_seen)}")
    date_html = " · ".join(date_parts)
    st.markdown(
        f"""
        <div class="narrative-card">
          <div class="narrative-card-header">
            <div class="narrative-card-title">{escape(name)}</div>
            {badge_html}
          </div>
          <div class="narrative-thesis">{escape(thesis)}</div>
          <div class="score-bar-outer">
            <div class="score-bar-inner" style="width:{score_pct}%;"></div>
          </div>
          <div class="narrative-footer">
            <span>Importance: {importance:.1f}/10</span>
            <span>·</span>
            <span>Mentions: {mentions}</span>
            {"<span>·</span><span>" + date_html + "</span>" if date_html else ""}
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_article_metadata(selected_article: pd.Series) -> None:
    st.markdown(
        f"""
        <div class="meta-grid">
          <div class="meta-card">
            <div class="meta-label">Source</div>
            <div class="meta-value">{escape(str(selected_article['source'] or 'Unknown'))}</div>
          </div>
          <div class="meta-card">
            <div class="meta-label">Published</div>
            <div class="meta-value">{escape(str(selected_article['published_date'] or 'Undated'))}</div>
          </div>
          <div class="meta-card">
            <div class="meta-label">Category</div>
            <div class="meta-value">{escape(str(selected_article['category'] or 'Uncategorized'))}</div>
          </div>
          <div class="meta-card">
            <div class="meta-label">Importance</div>
            <div class="meta-value">{escape(str(selected_article['importance_score']))} / 10</div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_summary_card(summary: str) -> None:
    st.markdown(
        f"""
        <div class="summary-card">
          <div class="summary-title">Analyst Summary</div>
          <div class="summary-body">{escape(summary or 'No summary available.')}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _js_literal(value: str) -> str:
    """Serialize a string for safe embedding inside inline component JavaScript."""

    return json.dumps(value).replace("</", "<\\/")


def _render_doc_preview(article_id: str, preview_text: str) -> None:
    storage_key = f"narrative_agent.article_notes.{article_id}"
    component_html = f"""
    <div class="annotator-shell">
      <div class="reader-toolbar">
        <div>
          <div class="reader-kicker">Source Text</div>
          <div class="reader-heading">Article Reader</div>
        </div>
        <div class="reader-actions">
          <button id="highlight-selection" type="button" title="Alt+H">Highlight</button>
          <button id="copy-notes" type="button">Copy Notes</button>
        </div>
      </div>
      <div class="reader-layout">
        <article id="source-reader" class="source-reader" tabindex="0" aria-label="Source document text"></article>
        <aside class="notes-panel" aria-label="Article notes">
          <div class="notes-panel-title">Notes</div>
          <div id="selected-text" class="selected-text">No selection</div>
          <textarea id="note-input" rows="4" placeholder="Take a note"></textarea>
          <div class="note-actions">
            <button id="save-note" type="button">Save Note</button>
            <button id="clear-draft" type="button">Clear</button>
          </div>
          <div id="note-status" class="note-status" role="status"></div>
          <div id="notes-list" class="notes-list"></div>
        </aside>
      </div>
    </div>
    <style>
      :root {{
        color-scheme: light;
      }}
      body {{
        margin: 0;
        background: transparent;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
        color: #172033;
      }}
      .annotator-shell {{
        background: rgba(248, 250, 252, 0.96);
        border: 1px solid rgba(23, 32, 51, 0.08);
        border-radius: 22px;
        padding: 0.9rem;
        box-sizing: border-box;
        height: 940px;
      }}
      .reader-toolbar {{
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 0.8rem;
        margin-bottom: 0.75rem;
      }}
      .reader-kicker {{
        color: #7c3f00;
        font-size: 0.72rem;
        font-weight: 700;
        letter-spacing: 0.14em;
        text-transform: uppercase;
      }}
      .reader-heading {{
        color: #172033;
        font-family: "Iowan Old Style", "Palatino Linotype", Georgia, serif;
        font-size: 1.25rem;
        font-weight: 700;
      }}
      .reader-actions,
      .note-actions {{
        display: flex;
        gap: 0.45rem;
        flex-wrap: wrap;
      }}
      button {{
        border: 1px solid rgba(23, 32, 51, 0.12);
        border-radius: 999px;
        background: #ffffff;
        color: #172033;
        cursor: pointer;
        font-weight: 700;
        padding: 0.48rem 0.76rem;
      }}
      button:hover {{
        border-color: rgba(29, 78, 216, 0.34);
        color: #1d4ed8;
      }}
      #save-note {{
        background: #172033;
        color: #ffffff;
        border-color: #172033;
      }}
      .reader-layout {{
        display: grid;
        grid-template-columns: minmax(0, 2.4fr) minmax(17rem, 0.58fr);
        gap: 0.85rem;
        height: calc(100% - 3.55rem);
      }}
      .source-reader {{
        background: #ffffff;
        border: 1px solid rgba(23, 32, 51, 0.08);
        border-radius: 18px;
        box-sizing: border-box;
        color: #172033;
        font-family: "SFMono-Regular", Menlo, Consolas, monospace;
        font-size: 1.02rem;
        line-height: 1.78;
        overflow: auto;
        padding: 1.25rem 1.35rem;
        white-space: pre-wrap;
        user-select: text;
      }}
      .source-reader:focus {{
        outline: 2px solid rgba(29, 78, 216, 0.25);
        outline-offset: 2px;
      }}
      mark.note-highlight {{
        background: rgba(250, 204, 21, 0.42);
        border-bottom: 2px solid rgba(202, 138, 4, 0.65);
        color: inherit;
        cursor: pointer;
        padding: 0.06rem 0.02rem;
      }}
      mark.note-highlight.is-active {{
        background: rgba(29, 78, 216, 0.18);
        border-bottom-color: #1d4ed8;
      }}
      mark.note-highlight.is-draft {{
        background: rgba(250, 204, 21, 0.58);
        border-bottom-color: #ca8a04;
      }}
      .notes-panel {{
        background: rgba(255, 255, 255, 0.88);
        border: 1px solid rgba(23, 32, 51, 0.08);
        border-radius: 18px;
        box-sizing: border-box;
        display: flex;
        flex-direction: column;
        min-height: 0;
        padding: 0.9rem;
      }}
      .notes-panel-title {{
        color: #7c3f00;
        font-size: 0.72rem;
        font-weight: 700;
        letter-spacing: 0.14em;
        margin-bottom: 0.55rem;
        text-transform: uppercase;
      }}
      .selected-text {{
        background: rgba(23, 32, 51, 0.05);
        border: 1px solid rgba(23, 32, 51, 0.08);
        border-radius: 14px;
        color: #5d6676;
        font-size: 0.86rem;
        line-height: 1.45;
        margin-bottom: 0.55rem;
        max-height: 5.6rem;
        overflow: auto;
        padding: 0.65rem;
      }}
      textarea {{
        background: #ffffff;
        border: 1px solid rgba(23, 32, 51, 0.12);
        border-radius: 14px;
        box-sizing: border-box;
        color: #172033;
        font: inherit;
        margin-bottom: 0.55rem;
        padding: 0.65rem;
        resize: vertical;
        width: 100%;
      }}
      textarea:focus {{
        border-color: rgba(29, 78, 216, 0.38);
        outline: 2px solid rgba(29, 78, 216, 0.14);
      }}
      .note-status {{
        color: #5d6676;
        font-size: 0.82rem;
        min-height: 1.2rem;
        padding: 0.35rem 0 0.15rem 0;
      }}
      .notes-list {{
        border-top: 1px solid rgba(23, 32, 51, 0.08);
        margin-top: 0.35rem;
        min-height: 0;
        overflow: auto;
        padding-top: 0.55rem;
      }}
      .note-card {{
        background: rgba(248, 250, 252, 0.96);
        border: 1px solid rgba(23, 32, 51, 0.08);
        border-radius: 14px;
        margin-bottom: 0.55rem;
        padding: 0.65rem;
      }}
      .note-card.active {{
        border-color: rgba(29, 78, 216, 0.34);
        box-shadow: 0 0 0 2px rgba(29, 78, 216, 0.10);
      }}
      .note-quote {{
        color: #172033;
        font-size: 0.86rem;
        font-weight: 700;
        line-height: 1.4;
        margin-bottom: 0.35rem;
      }}
      .note-body {{
        color: #5d6676;
        font-size: 0.86rem;
        line-height: 1.45;
        white-space: pre-wrap;
      }}
      .note-meta {{
        align-items: center;
        display: flex;
        justify-content: space-between;
        gap: 0.5rem;
        margin-top: 0.45rem;
      }}
      .note-date {{
        color: #6a7485;
        font-size: 0.72rem;
      }}
      .delete-note {{
        color: #b91c1c;
        font-size: 0.76rem;
        padding: 0.3rem 0.5rem;
      }}
      .empty-note {{
        border: 1px dashed rgba(23, 32, 51, 0.18);
        border-radius: 14px;
        color: #6a7485;
        font-size: 0.86rem;
        padding: 0.75rem;
      }}
      @media (max-width: 760px) {{
        .annotator-shell {{
          height: 980px;
        }}
        .reader-layout {{
          grid-template-columns: 1fr;
          grid-template-rows: 1fr minmax(20rem, 0.8fr);
        }}
      }}
    </style>
    <script>
      const sourceText = {_js_literal(preview_text)};
      const storageKey = {_js_literal(storage_key)};
      const reader = document.getElementById("source-reader");
      const noteInput = document.getElementById("note-input");
      const selectedText = document.getElementById("selected-text");
      const statusBox = document.getElementById("note-status");
      const notesList = document.getElementById("notes-list");
      const highlightButton = document.getElementById("highlight-selection");
      const saveButton = document.getElementById("save-note");
      const clearButton = document.getElementById("clear-draft");
      const copyButton = document.getElementById("copy-notes");
      let draft = null;
      let activeNoteId = null;
      let storageAvailable = true;
      let notes = loadNotes();

      function escapeHtml(value) {{
        return String(value)
          .replaceAll("&", "&amp;")
          .replaceAll("<", "&lt;")
          .replaceAll(">", "&gt;")
          .replaceAll('"', "&quot;")
          .replaceAll("'", "&#039;");
      }}

      function makeId() {{
        if (window.crypto && window.crypto.randomUUID) {{
          return window.crypto.randomUUID();
        }}
        return String(Date.now()) + "-" + String(Math.random()).slice(2);
      }}

      function loadNotes() {{
        try {{
          const parsed = JSON.parse(window.localStorage.getItem(storageKey) || "[]");
          storageAvailable = true;
          if (!Array.isArray(parsed)) {{
            return [];
          }}
          return parsed.filter(isValidNote);
        }} catch (error) {{
          storageAvailable = false;
          return [];
        }}
      }}

      function isValidNote(note) {{
        return note
          && Number.isInteger(note.start)
          && Number.isInteger(note.end)
          && note.start >= 0
          && note.end > note.start
          && note.end <= sourceText.length
          && typeof note.text === "string";
      }}

      function persistNotes() {{
        try {{
          window.localStorage.setItem(storageKey, JSON.stringify(notes));
          storageAvailable = true;
          return true;
        }} catch (error) {{
          storageAvailable = false;
          return false;
        }}
      }}

      function setStatus(message) {{
        statusBox.textContent = message || "";
      }}

      function renderReader() {{
        const ranges = notes
          .filter(isValidNote)
          .map((note) => ({{ ...note, kind: "note" }}));
        if (draft) {{
          ranges.push({{ ...draft, id: "draft", kind: "draft" }});
        }}
        const sorted = ranges.sort((a, b) => a.start - b.start || a.end - b.end);
        let cursor = 0;
        let html = "";
        for (const range of sorted) {{
          if (range.start < cursor) {{
            continue;
          }}
          html += escapeHtml(sourceText.slice(cursor, range.start));
          const highlightedText = sourceText.slice(range.start, range.end);
          const classNames = ["note-highlight"];
          if (range.kind === "draft") {{
            classNames.push("is-draft");
          }} else if (range.id === activeNoteId) {{
            classNames.push("is-active");
          }}
          const markerAttribute = range.kind === "note"
            ? ` data-note-id="${{escapeHtml(range.id)}}"`
            : ' data-draft-highlight="true"';
          html += `<mark class="${{classNames.join(" ")}}"${{markerAttribute}}>${{escapeHtml(highlightedText)}}</mark>`;
          cursor = range.end;
        }}
        html += escapeHtml(sourceText.slice(cursor));
        reader.innerHTML = html || '<span class="empty-note">Preview unavailable.</span>';
        renderNotes();
      }}

      function renderNotes() {{
        if (!notes.length) {{
          notesList.innerHTML = '<div class="empty-note">No notes yet</div>';
          return;
        }}
        const sorted = [...notes].sort((a, b) => (b.createdAt || "").localeCompare(a.createdAt || ""));
        notesList.innerHTML = sorted.map((note) => {{
          const activeClass = note.id === activeNoteId ? " active" : "";
          const dateLabel = note.createdAt ? new Date(note.createdAt).toLocaleString() : "";
          return `
            <div class="note-card${{activeClass}}" data-note-card-id="${{escapeHtml(note.id)}}">
              <div class="note-quote">${{escapeHtml(note.text)}}</div>
              <div class="note-body">${{escapeHtml(note.note || "No note text")}}</div>
              <div class="note-meta">
                <span class="note-date">${{escapeHtml(dateLabel)}}</span>
                <button class="delete-note" type="button" data-delete-id="${{escapeHtml(note.id)}}">Delete</button>
              </div>
            </div>
          `;
        }}).join("");
      }}

      function selectionOffsets() {{
        const selection = window.getSelection();
        if (!selection || selection.rangeCount === 0) {{
          return null;
        }}
        const selected = selection.toString();
        if (!selected.trim()) {{
          return null;
        }}
        const range = selection.getRangeAt(0);
        if (!reader.contains(range.commonAncestorContainer)) {{
          return null;
        }}
        const preRange = document.createRange();
        preRange.selectNodeContents(reader);
        preRange.setEnd(range.startContainer, range.startOffset);
        const start = preRange.toString().length;
        const end = start + selected.length;
        if (start < 0 || end > sourceText.length || end <= start) {{
          return null;
        }}
        return {{ start, end, text: sourceText.slice(start, end) }};
      }}

      function captureSelection(options = {{}}) {{
        const focusNote = options.focusNote !== false;
        const quiet = options.quiet === true;
        const nextDraft = selectionOffsets();
        if (!nextDraft) {{
          if (!quiet) {{
            setStatus(draft ? "Selection ready." : "Select source text first.");
          }}
          if (draft && focusNote) {{
            noteInput.focus();
          }}
          return false;
        }}
        draft = nextDraft;
        selectedText.textContent = draft.text;
        noteInput.value = "";
        renderReader();
        if (!quiet) {{
          setStatus("Selection ready.");
        }}
        if (focusNote) {{
          noteInput.focus();
        }}
        return true;
      }}

      function clearDraft() {{
        draft = null;
        selectedText.textContent = "No selection";
        noteInput.value = "";
        setStatus("");
        window.getSelection()?.removeAllRanges();
        renderReader();
      }}

      function saveDraft() {{
        if (!draft && !captureSelection({{ focusNote: false, quiet: true }})) {{
          setStatus("No selection captured.");
          return;
        }}
        const noteText = noteInput.value.trim();
        const note = {{
          id: makeId(),
          start: draft.start,
          end: draft.end,
          text: draft.text,
          note: noteText,
          createdAt: new Date().toISOString(),
        }};
        notes.push(note);
        activeNoteId = note.id;
        draft = null;
        selectedText.textContent = "No selection";
        noteInput.value = "";
        window.getSelection()?.removeAllRanges();
        renderReader();
        const persisted = persistNotes();
        setStatus(persisted ? "Note saved." : "Note saved for this session. Browser storage unavailable.");
      }}

      function activateNote(noteId) {{
        activeNoteId = noteId;
        renderReader();
        const mark = [...reader.querySelectorAll("[data-note-id]")]
          .find((element) => element.dataset.noteId === noteId);
        if (mark) {{
          mark.scrollIntoView({{ behavior: "smooth", block: "center" }});
        }}
      }}

      function deleteNote(noteId) {{
        notes = notes.filter((note) => note.id !== noteId);
        if (activeNoteId === noteId) {{
          activeNoteId = null;
        }}
        renderReader();
        const persisted = persistNotes();
        setStatus(persisted ? "Note deleted." : "Note deleted for this session. Browser storage unavailable.");
      }}

      async function copyNotes() {{
        const payload = notes.map((note) => ({{
          selected_text: note.text,
          note: note.note || "",
          created_at: note.createdAt,
        }}));
        try {{
          await navigator.clipboard.writeText(JSON.stringify(payload, null, 2));
          setStatus("Notes copied.");
        }} catch (error) {{
          setStatus("Copy unavailable in this browser.");
        }}
      }}

      function isHighlightShortcut(event) {{
        return event.altKey && (event.code === "KeyH" || String(event.key || "").toLowerCase() === "h");
      }}

      function focusDraftNote() {{
        if (draft) {{
          setStatus("Selection ready.");
          noteInput.focus();
          return true;
        }}
        return captureSelection({{ focusNote: true, quiet: false }});
      }}

      document.addEventListener("keydown", (event) => {{
        if (isHighlightShortcut(event)) {{
          event.preventDefault();
          focusDraftNote();
        }}
      }});
      reader.addEventListener("mouseup", () => {{
        window.setTimeout(() => captureSelection({{ focusNote: false, quiet: true }}), 0);
      }});
      reader.addEventListener("keyup", (event) => {{
        if (event.shiftKey || event.metaKey) {{
          window.setTimeout(() => captureSelection({{ focusNote: false, quiet: true }}), 0);
        }}
      }});
      highlightButton.addEventListener("click", focusDraftNote);
      saveButton.addEventListener("click", saveDraft);
      clearButton.addEventListener("click", clearDraft);
      copyButton.addEventListener("click", copyNotes);
      reader.addEventListener("click", (event) => {{
        const mark = event.target.closest("[data-note-id]");
        if (mark) {{
          activateNote(mark.dataset.noteId);
        }}
      }});
      notesList.addEventListener("click", (event) => {{
        const deleteButton = event.target.closest("[data-delete-id]");
        if (deleteButton) {{
          deleteNote(deleteButton.dataset.deleteId);
          return;
        }}
        const card = event.target.closest("[data-note-card-id]");
        if (card) {{
          activateNote(card.dataset.noteCardId);
        }}
      }});

      renderReader();
    </script>
    """
    components.html(component_html, height=960, scrolling=False)


def _render_media_shell_intro(image_count: int, preview_text: str, truncated: bool) -> None:
    chips = [
        f'<span class="media-chip">Preview length: {len(preview_text):,} chars</span>',
        f'<span class="media-chip">Images: {image_count}</span>',
    ]
    if truncated:
        chips.append('<span class="media-chip">Text preview truncated</span>')
    st.markdown(
        f"""
        <div class="media-shell">
          <div class="media-intro">
            <div>
              <div class="media-title">Document And Media Board</div>
              <div class="media-note">
                Read the extracted source text and inspect embedded visuals in one place instead of jumping between separate blocks.
              </div>
            </div>
          </div>
          <div class="media-chip-row">{''.join(chips)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _render_image_rail(preview_images: list[bytes], docx_missing: bool) -> None:
    st.markdown(
        """
        <div class="image-rail">
          <div class="image-rail-title">Embedded Visuals</div>
          <div class="image-rail-copy">Images pulled from the original DOCX are pinned beside the text preview for easier comparison.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if docx_missing:
        st.markdown(
            '<div class="image-empty">Image preview unavailable because the source DOCX could not be found.</div>',
            unsafe_allow_html=True,
        )
        return

    if not preview_images:
        st.markdown(
            '<div class="image-empty">No embedded images were found in this DOCX.</div>',
            unsafe_allow_html=True,
        )
        return

    image_cols = st.columns(min(len(preview_images), DOCX_IMAGE_LIMIT), gap="small")
    for index, image_payload in enumerate(preview_images[:DOCX_IMAGE_LIMIT], start=1):
        with image_cols[index - 1]:
            caption = "Lead embedded image" if index == 1 else f"Embedded image {index}"
            st.image(image_payload, caption=caption, width="stretch")

    if len(preview_images) == DOCX_IMAGE_LIMIT:
        st.caption("Image preview limited for display.")


def _render_graph_legend() -> None:
    legend_items = [
        (NODE_COLORS["article"], "Article", NODE_SHAPES["article"]),
        (NODE_COLORS["entity"], "Entity", NODE_SHAPES["entity"]),
        (NODE_COLORS["theme"], "Theme", NODE_SHAPES["theme"]),
        (NODE_COLORS["narrative"], "Narrative", NODE_SHAPES["narrative"]),
        ("#16a34a", "New (compare)", "dot"),
        ("#dc2626", "Removed (compare)", "dot"),
    ]
    chips = "".join(
        f'<span class="legend-chip">'
        f'<span class="legend-dot" style="background:{color};flex-shrink:0;"></span>'
        f'{escape(label)}'
        f'</span>'
        for color, label, _ in legend_items
    )
    st.markdown(
        f'<div class="legend-strip">{chips}</div>',
        unsafe_allow_html=True,
    )
    st.caption("Larger nodes = higher importance or stronger article linkage. Hover nodes/edges for details.")


def _render_chart_card(title: str, chart: alt.Chart) -> None:
    st.markdown(
        f"""
        <div class="overview-card">
          <div class="section-kicker">Distribution</div>
          <div class="section-title" style="font-size:1.2rem;">{escape(title)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.altair_chart(chart, width="stretch")


def _top_count_frame(dataframe: pd.DataFrame, column_name: str, label: str, limit: int = 8) -> pd.DataFrame:
    counts = (
        dataframe[column_name]
        .fillna("")
        .replace("", "Unspecified")
        .value_counts()
        .head(limit)
        .rename_axis(label)
        .reset_index(name="count")
    )
    return counts


def _count_chart(dataframe: pd.DataFrame, category_column: str, label: str, color: str) -> alt.Chart:
    chart_data = _top_count_frame(dataframe, category_column, label)
    return (
        alt.Chart(chart_data)
        .mark_bar(cornerRadiusTopRight=7, cornerRadiusBottomRight=7, color=color)
        .encode(
            x=alt.X("count:Q", title="Articles"),
            y=alt.Y(f"{label}:N", sort="-x", title=None),
            tooltip=[alt.Tooltip(f"{label}:N", title=label.title()), alt.Tooltip("count:Q", title="Articles")],
        )
        .properties(height=280)
    )


def _render_overview_tab(
    filtered_articles: pd.DataFrame,
    filtered_narratives: pd.DataFrame,
    selected_week: str,
) -> None:
    if filtered_articles.empty:
        st.info("No articles match the current filters.")
        return

    avg_importance = float(filtered_articles["importance_score"].fillna(0).mean())
    source_count = int(filtered_articles["source"].replace("", pd.NA).dropna().nunique())
    narrative_count = int(filtered_narratives["name"].nunique()) if not filtered_narratives.empty else 0
    high_priority_count = int((filtered_articles["importance_score"].fillna(0) >= 8).sum())

    top_row = st.columns(4)
    with top_row[0]:
        _render_metric_card("Articles in Scope", str(len(filtered_articles)), "Current filtered working set")
    with top_row[1]:
        _render_metric_card(
            "Average Importance",
            f"{avg_importance:.1f}",
            "Mean article score across the current view",
        )
    with top_row[2]:
        _render_metric_card("Sources Active", str(source_count), "Distinct publications in view")
    with top_row[3]:
        _render_metric_card("Priority Coverage", str(high_priority_count), "Articles scoring 8 or higher")

    st.markdown("")
    left_col, right_col = st.columns([1.15, 0.85], gap="large")

    with left_col:
        _render_section_intro(
            "Coverage Shape",
            "Quick distribution checks help reveal whether the week is concentrated in one source or one topic cluster.",
            kicker="Overview",
        )
        chart_cols = st.columns(2, gap="large")
        with chart_cols[0]:
            _render_chart_card("Top Sources", _count_chart(filtered_articles, "source", "source", "#1d4ed8"))
        with chart_cols[1]:
            _render_chart_card("Top Categories", _count_chart(filtered_articles, "category", "category", "#0f766e"))

    with right_col:
        _render_section_intro(
            "Operating Snapshot",
            "This panel compresses the current filtered view into the main signals you would scan before reading individual documents.",
            kicker="Status",
        )
        recent_view = filtered_articles[
            ["title", "source", "published_date", "importance_score"]
        ].head(8).rename(columns={"published_date": "date", "importance_score": "importance"})
        st.dataframe(recent_view, width="stretch", hide_index=True, height=318)
        st.markdown(
            f"""
            <div class="panel-card" style="margin-top:0.8rem;">
              <div class="section-kicker">Narratives</div>
              <div class="section-title" style="font-size:1.25rem;">{narrative_count} active narrative tracks</div>
              <div class="section-copy">Week filter: {escape(selected_week)}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )


def _render_graph_chat_result(result: dict[str, Any]) -> None:
    facts = result.get("facts", [])
    citations = result.get("citations", [])
    evaluation = result.get("evaluation", {})
    warnings = result.get("warnings", [])

    if warnings:
        for warning in warnings:
            st.warning(str(warning))

    evidence_rows = [
        {
            "type": fact.get("kind", ""),
            "statement": fact.get("statement", ""),
            "relationship": (
                f"{fact.get('source_node')} --{fact.get('relationship')}--> {fact.get('target_node')}"
                if fact.get("source_node") and fact.get("relationship") and fact.get("target_node")
                else ""
            ),
            "confidence": fact.get("confidence", ""),
            "article": fact.get("title", ""),
            "date": fact.get("published_date", ""),
        }
        for fact in facts
    ]
    citation_rows = [
        {
            "article": citation.get("title", ""),
            "date": citation.get("published_date", ""),
            "source": citation.get("source", ""),
            "evidence": citation.get("quote", ""),
        }
        for citation in citations
    ]

    if evidence_rows:
        with st.expander("Retrieved graph facts", expanded=False):
            st.dataframe(pd.DataFrame(evidence_rows), width="stretch", hide_index=True)

    if citation_rows:
        with st.expander("Citations", expanded=False):
            st.dataframe(pd.DataFrame(citation_rows), width="stretch", hide_index=True)

    if evaluation:
        st.caption(
            "Evaluation: "
            f"{evaluation.get('fact_count', 0)} facts, "
            f"{evaluation.get('edge_fact_count', 0)} edge facts, "
            f"{evaluation.get('citation_count', 0)} citations, "
            f"LLM used: {evaluation.get('llm_used', False)}"
        )


def _render_chat_text(content: str) -> None:
    """Render chat text without treating currency values as Markdown math."""

    st.markdown(str(content).replace("$", r"\$"))


def _render_ask_graph_tab(
    selected_week: str,
    default_backend: str = "Neo4j",
    *,
    neo4j_available: bool = True,
) -> None:
    _render_section_intro(
        "Ask The Graph",
        "Ask for event chains, relationship logic, narratives, and article evidence with graph-backed citations.",
        kicker="GraphRAG",
    )
    chat_week = None if selected_week == "All" else selected_week
    control_left, control_middle, control_right = st.columns([0.40, 0.30, 0.30], gap="large")
    with control_left:
        st.caption(f"Scope: {chat_week or 'All indexed weeks'}")
    with control_middle:
        if not neo4j_available:
            st.info("Graph chat is unavailable because Neo4j is not reachable.")
            return
        backend_options = ["Neo4j"]
        graph_backend = st.selectbox(
            "GraphRAG backend",
            options=backend_options,
            index=backend_options.index(default_backend) if default_backend in backend_options else 0,
            key="graph_rag_backend_source",
        )
    with control_right:
        deep_immediately = st.toggle(
            "Deep answer immediately",
            value=False,
            key="graph_rag_deep_now",
        )

    if "graph_chat_messages" not in st.session_state:
        st.session_state.graph_chat_messages = []

    if st.button("Clear chat", key="graph_chat_clear"):
        st.session_state.graph_chat_messages = []

    for message_index, message in enumerate(st.session_state.graph_chat_messages):
        with st.chat_message(message["role"]):
            _render_chat_text(message["content"])
            if message.get("result"):
                _render_graph_chat_result(message["result"])
                evaluation = message["result"].get("evaluation", {})
                can_deepen = (
                    message.get("role") == "assistant"
                    and message.get("backend") == "neo4j"
                    and message.get("question")
                    and not evaluation.get("deep_search", False)
                )
                if can_deepen and st.button(
                    "Deep answer with embeddings",
                    key=f"graph_chat_deep_{message_index}",
                ):
                    with st.spinner("Running embedding search and LLM synthesis..."):
                        try:
                            deep_result = answer_question(
                                str(message["question"]),
                                week=message.get("week"),
                                use_llm=True,
                                deep_search=True,
                                backend="neo4j",
                            )
                        except Exception as exc:
                            st.session_state.graph_chat_messages.append(
                                {
                                    "role": "assistant",
                                    "content": f"Deep answer failed: {exc}",
                                }
                            )
                            st.rerun()
                            return
                    st.session_state.graph_chat_messages.append(
                        {
                            "role": "assistant",
                            "content": deep_result.answer,
                            "result": deep_result.model_dump(),
                            "question": message["question"],
                            "week": message.get("week"),
                            "backend": "neo4j",
                        }
                    )
                    st.rerun()
                    return

    with st.form("graph_rag_query_form", clear_on_submit=True):
        prompt = st.text_area(
            "Question",
            placeholder="Ask a follow-up about entities, relationships, valuation, or recent changes",
            key="graph_rag_question",
            height=96,
        )
        button_label = "Ask follow-up" if st.session_state.graph_chat_messages else "Ask graph"
        submitted = st.form_submit_button(button_label, type="primary")

    if not submitted:
        return

    prompt = prompt.strip()
    if not prompt:
        st.warning("Please enter a question.")
        return

    st.session_state.graph_chat_messages.append({"role": "user", "content": prompt})
    with st.spinner("Querying graph facts..."):
        try:
            result = answer_question(
                prompt,
                week=chat_week,
                use_llm=deep_immediately,
                deep_search=deep_immediately,
                backend="neo4j",
            )
        except Exception as exc:
            backend_name = graph_backend
            st.session_state.graph_chat_messages.append(
                {
                    "role": "assistant",
                    "content": f"{backend_name} backend is not available yet: {exc}",
                }
            )
            st.rerun()
            return
    st.session_state.graph_chat_messages.append(
        {
            "role": "assistant",
            "content": result.answer,
            "result": result.model_dump(),
            "question": prompt,
            "week": chat_week,
            "backend": "neo4j",
        }
    )
    st.rerun()


def main() -> None:
    st.set_page_config(
        page_title="Narrative Agent Dashboard",
        page_icon="🗺️",
        layout="wide",
        initial_sidebar_state="expanded",
    )
    inject_dashboard_css()

    registry_version = _path_version(PROCESSED_MANIFEST_PATH)
    data_backend = "Processed JSON"
    neo4j_warning = ""

    if neo4j_dashboard_available(0):
        try:
            metrics = load_metrics_neo4j(0)
            articles_df = load_articles_neo4j(0)
            narratives_df = load_narratives_neo4j(0)
            entities_df = load_entities_neo4j(0)
            data_backend = "Neo4j"
        except Exception as exc:
            neo4j_warning = f"Neo4j dashboard read failed; using processed JSON fallback. Details: {exc}"
            data_backend = "Processed JSON"

    if data_backend == "Processed JSON":
        metrics = load_metrics(registry_version)
        articles_df = load_articles(registry_version)
        narratives_df = load_narratives(registry_version)
        entities_df = load_entities(registry_version)
        if articles_df.empty:
            st.error("Neo4j is unavailable and no processed extraction artifacts were found. Run the pipeline first.")
            if neo4j_warning:
                st.caption(neo4j_warning)
            return

    st.sidebar.markdown("## Filters")
    st.sidebar.caption("Scope the intelligence view before drilling into articles, narratives, and graph relationships.")
    st.sidebar.caption(f"Primary graph source: {data_backend}")
    if neo4j_warning:
        st.sidebar.warning(neo4j_warning)
    week_options = ["All", *sorted([value for value in articles_df["week_label"].dropna().unique()], reverse=True)]
    source_options = ["All", *sorted([value for value in articles_df["source"].dropna().unique() if value])]
    category_options = ["All", *sorted([value for value in articles_df["category"].dropna().unique() if value])]

    selected_week = st.sidebar.selectbox("Week", options=week_options)
    selected_source = st.sidebar.selectbox("Source", options=source_options)
    selected_category = st.sidebar.selectbox("Category", options=category_options)
    search_term = st.sidebar.text_input("Search", value="")

    st.sidebar.markdown("---")
    st.sidebar.markdown(
        f"""
        <div class="sidebar-db-card">
          <div class="sidebar-db-kicker">Database Snapshot</div>
          <div class="sidebar-stat-row">
            <span class="sidebar-stat-label">Articles</span>
            <span class="sidebar-stat-value">{metrics['articles']}</span>
          </div>
          <div class="sidebar-stat-row">
            <span class="sidebar-stat-label">Entities</span>
            <span class="sidebar-stat-value">{metrics['entities']}</span>
          </div>
          <div class="sidebar-stat-row">
            <span class="sidebar-stat-label">Themes</span>
            <span class="sidebar-stat-value">{metrics['themes']}</span>
          </div>
          <div class="sidebar-stat-row">
            <span class="sidebar-stat-label">Narratives</span>
            <span class="sidebar-stat-value">{metrics['narratives']}</span>
          </div>
          <div class="sidebar-stat-row">
            <span class="sidebar-stat-label">Graph Edges</span>
            <span class="sidebar-stat-value">{metrics['graph_edges']}</span>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    filtered_articles = _filter_articles(
        articles_df,
        selected_week,
        selected_source,
        selected_category,
        search_term,
    )
    filtered_article_ids = set(filtered_articles["id"].astype(str))
    filtered_narratives = _filter_related_rows(narratives_df, filtered_article_ids)
    active_filters = [
        f"Week: {selected_week}",
        f"Source: {selected_source}",
        f"Category: {selected_category}",
        f"Search: {search_term or 'None'}",
    ]
    hero_badges = "".join(f'<span class="hero-badge">{escape(item)}</span>' for item in active_filters)
    total_filtered = len(filtered_articles)
    avg_importance = float(filtered_articles["importance_score"].fillna(0).mean()) if total_filtered else 0.0

    st.markdown(
        f"""
        <div class="hero-panel">
          <div class="hero-kicker">Narrative Intelligence Console</div>
          <div class="hero-title">Narrative Agent Dashboard</div>
          <div class="hero-copy">
            Structured weekly coverage, entity tracking, narrative movement, and graph relationships in one operating view.
            Current scope includes {total_filtered} article(s) with an average importance score of {avg_importance:.1f}.
          </div>
          <div class="hero-badges">{hero_badges}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    metric_columns = st.columns(5)
    with metric_columns[0]:
        _render_metric_card("Articles", str(metrics["articles"]), "Total article records")
    with metric_columns[1]:
        _render_metric_card("Entities", str(metrics["entities"]), "Distinct tracked entities")
    with metric_columns[2]:
        _render_metric_card("Themes", str(metrics["themes"]), "Reusable topic clusters")
    with metric_columns[3]:
        _render_metric_card("Narratives", str(metrics["narratives"]), "Trackable narrative lines")
    with metric_columns[4]:
        _render_metric_card("Graph Edges", str(metrics["graph_edges"]), "Relationship statements")

    ask_tab, overview_tab, articles_tab, narratives_tab, entities_tab, trends_tab, graph_tab = st.tabs(
        ["Ask Graph", "Overview", "Articles", "Narratives", "Entities", "Narrative Trends", "Graph View"]
    )

    with ask_tab:
        _render_ask_graph_tab(
            selected_week,
            default_backend=data_backend,
            neo4j_available=data_backend == "Neo4j",
        )

    with overview_tab:
        _render_overview_tab(filtered_articles, filtered_narratives, selected_week)

    with articles_tab:
        if filtered_articles.empty:
            st.info("No articles match the current filters.")
        else:
            _render_section_intro(
                "Article Workbench",
                "Use this pane to inspect the structured extraction before opening the full source reader below.",
                kicker="Primary View",
            )
            article_options = filtered_articles["selection_label"].tolist()

            st.caption(f"{len(filtered_articles)} article(s) match the current filter set.")
            selected_label = st.selectbox("Select article", article_options, key="article_select")
            selected_article = filtered_articles[
                filtered_articles["selection_label"] == selected_label
            ].iloc[0]

            workbench_left, workbench_right = st.columns([1.15, 0.85], gap="large")
            with workbench_left:
                st.markdown(f"## {selected_article['title']}")
                _render_article_metadata(selected_article)
                _render_summary_card(str(selected_article["summary"] or ""))

            with workbench_right:
                entity_names = _split_names(selected_article["linked_entities"])
                theme_names = _split_names(selected_article["linked_themes"])
                narrative_names = _split_names(selected_article["linked_narratives"])
                _render_signal_group("Entities", entity_names, "No linked entities")
                _render_signal_group("Themes", theme_names, "No linked themes")
                _render_signal_group("Narratives", narrative_names, "No linked narratives")

            st.markdown('<div class="article-source-divider"></div>', unsafe_allow_html=True)
            original_file_path = str(selected_article["original_file_path"] or "")
            preview_text = preview_docx(original_file_path)
            docx_missing = preview_text == "Original DOCX not found"
            preview_images = preview_docx_images(original_file_path)
            preview_truncated = False if docx_missing else _is_docx_preview_truncated(original_file_path)
            source_header_left, source_header_right = st.columns([1.08, 0.92], gap="large")
            with source_header_left:
                _render_section_intro(
                    "Source Document",
                    "Read the original document in a larger source panel with article-level notes.",
                    kicker="Document",
                )
            with source_header_right:
                _render_media_shell_intro(len(preview_images), preview_text, preview_truncated)
            _render_image_rail(preview_images, docx_missing)
            _render_doc_preview(str(selected_article["id"]), preview_text)
            if docx_missing:
                st.info("Original DOCX not found")
            elif preview_truncated:
                st.caption("Preview truncated for display.")

    with narratives_tab:
        _render_section_intro(
            "Narrative Register",
            "Filter the narrative layer independently to review theses, status labels, and mention longevity.",
            kicker="Narratives",
        )
        status_options = [
            "emerging",
            "strengthening",
            "weakening",
            "transitioning",
            "reversed",
            "stable",
        ]
        selected_statuses = st.multiselect(
            "Narrative status",
            options=status_options,
            default=status_options,
        )

        filtered_narratives = filtered_narratives[
            filtered_narratives["status"].isin(selected_statuses)
        ]
        if search_term:
            filtered_narratives = filtered_narratives[
                filtered_narratives["name"].str.contains(search_term, case=False, na=False)
                | filtered_narratives["thesis"].str.contains(search_term, case=False, na=False)
            ]

        narrative_metric_cols = st.columns(3)
        with narrative_metric_cols[0]:
            _render_metric_card("Visible Narratives", str(len(filtered_narratives)), "Rows after filters")
        with narrative_metric_cols[1]:
            _render_metric_card(
                "Average Importance",
                f"{filtered_narratives['importance_score'].fillna(0).mean():.1f}" if not filtered_narratives.empty else "0.0",
                "Mean importance across visible narratives",
            )
        with narrative_metric_cols[2]:
            _render_metric_card(
                "Average Mentions",
                f"{filtered_narratives['mention_count'].fillna(0).mean():.1f}" if not filtered_narratives.empty else "0.0",
                "Mean article mentions per narrative",
            )

        st.markdown("")
        if filtered_narratives.empty:
            st.info("No narratives match the current filters.")
        else:
            CARD_LIMIT = 12
            display_rows = filtered_narratives.head(CARD_LIMIT)
            for _, row in display_rows.iterrows():
                _render_narrative_card(row)
            if len(filtered_narratives) > CARD_LIMIT:
                with st.expander(f"Show all {len(filtered_narratives)} narratives as table"):
                    st.dataframe(
                        filtered_narratives[
                            [
                                "name",
                                "thesis",
                                "status",
                                "importance_score",
                                "mention_count",
                                "first_seen_date",
                                "last_seen_date",
                                "linked_articles",
                            ]
                        ],
                        width="stretch",
                        hide_index=True,
                    )

    with entities_tab:
        _render_section_intro(
            "Entity Register",
            "Review the monitored actors and the article/narrative contexts in which they appear.",
            kicker="Entities",
        )
        filtered_entities = _filter_related_rows(entities_df, filtered_article_ids)
        if search_term:
            filtered_entities = filtered_entities[
                filtered_entities["name"].str.contains(search_term, case=False, na=False)
                | filtered_entities["description"].str.contains(search_term, case=False, na=False)
            ]

        entity_metric_cols = st.columns(3)
        with entity_metric_cols[0]:
            _render_metric_card("Visible Entities", str(len(filtered_entities)), "Rows after filters")
        with entity_metric_cols[1]:
            _render_metric_card(
                "Entity Types",
                str(filtered_entities["type"].replace("", pd.NA).dropna().nunique()),
                "Distinct entity classes in view",
            )
        with entity_metric_cols[2]:
            relationship_count = int(filtered_entities["graph_relationships"].fillna("").astype(str).str.len().gt(0).sum())
            _render_metric_card("Linked Relationships", str(relationship_count), "Entities with graph evidence")

        st.dataframe(
            filtered_entities[["name", "type", "description", "linked_articles", "graph_relationships"]],
            width="stretch",
            hide_index=True,
        )

    with trends_tab:
        _render_section_intro(
            "Narrative Movement",
            "Trend labels are computed from article-linked appearances inside the selected week window.",
            kicker="Trends",
        )
        if selected_week == "All":
            st.info("Select a specific week in the sidebar to view narrative trends.")
        else:
            trend_warning = ""
            if data_backend == "Neo4j":
                try:
                    trend_df = load_narrative_trends_neo4j(selected_week, 0)
                except Exception as exc:
                    trend_warning = f"Neo4j narrative trend read failed; showing processed JSON fallback. Details: {exc}"
                    trend_df = load_narrative_trends(selected_week, registry_version)
            elif data_backend == "Processed JSON":
                trend_df = load_narrative_trends(selected_week, registry_version)
            else:
                st.info("Narrative trend data is unavailable because Neo4j could not be read.")
                return

            if trend_warning:
                st.warning(trend_warning)
            if search_term:
                trend_df = trend_df[
                    trend_df["name"].str.contains(search_term, case=False, na=False)
                    | trend_df["thesis"].str.contains(search_term, case=False, na=False)
                ]

            status_options = ["NEW", "STRENGTHENING", "RECURRING", "WEAKENING"]
            selected_trend_statuses = st.multiselect(
                "Trend status",
                options=status_options,
                default=status_options,
                key="trend_status_filter",
            )
            trend_df = trend_df[trend_df["trend_status"].isin(selected_trend_statuses)]
            trend_df = trend_df.sort_values(by=["mention_count", "name"], ascending=[False, True])

            trend_metric_cols = st.columns(4)
            with trend_metric_cols[0]:
                _render_metric_card("Visible Trends", str(len(trend_df)), "Rows after filters")
            with trend_metric_cols[1]:
                strongest_count = int((trend_df["trend_status"] == "STRENGTHENING").sum())
                _render_metric_card("Strengthening", str(strongest_count), "Narratives accelerating this week")
            with trend_metric_cols[2]:
                new_count = int((trend_df["trend_status"] == "NEW").sum())
                _render_metric_card("New Signals", str(new_count), "Narratives first seen this week")
            with trend_metric_cols[3]:
                weakening_count = int((trend_df["trend_status"] == "WEAKENING").sum())
                _render_metric_card("Weakening", str(weakening_count), "Narratives losing momentum")

            st.markdown("")
            if not trend_df.empty:
                # Status distribution bar chart
                status_counts = (
                    trend_df["trend_status"]
                    .value_counts()
                    .rename_axis("status")
                    .reset_index(name="count")
                )
                status_color_map = {
                    "NEW": "#0f766e",
                    "STRENGTHENING": "#1d4ed8",
                    "RECURRING": "#6b7280",
                    "WEAKENING": "#b45309",
                }
                status_counts["color"] = status_counts["status"].map(
                    lambda s: status_color_map.get(s, "#556277")
                )
                trend_chart = (
                    alt.Chart(status_counts)
                    .mark_bar(cornerRadiusTopRight=7, cornerRadiusBottomRight=7)
                    .encode(
                        x=alt.X("count:Q", title="Count"),
                        y=alt.Y("status:N", sort="-x", title=None),
                        color=alt.Color(
                            "status:N",
                            scale=alt.Scale(
                                domain=list(status_color_map.keys()),
                                range=list(status_color_map.values()),
                            ),
                            legend=None,
                        ),
                        tooltip=[
                            alt.Tooltip("status:N", title="Status"),
                            alt.Tooltip("count:Q", title="Count"),
                        ],
                    )
                    .properties(height=160, title="Trends by Status")
                )
                st.altair_chart(trend_chart, width="stretch")

                # Styled table with inline status badges rendered as text column
                display_df = trend_df[
                    [
                        "name",
                        "trend_status",
                        "mention_count",
                        "first_seen_date",
                        "last_seen_date",
                    ]
                ].rename(columns={"trend_status": "status"})
                st.dataframe(
                    display_df,
                    width="stretch",
                    hide_index=True,
                    column_config={
                        "name": st.column_config.TextColumn("Narrative", width="large"),
                        "status": st.column_config.TextColumn("Status"),
                        "mention_count": st.column_config.NumberColumn("Mentions", format="%d"),
                        "first_seen_date": st.column_config.TextColumn("First Seen"),
                        "last_seen_date": st.column_config.TextColumn("Last Seen"),
                    },
                )
            else:
                st.info("No trends match the selected filters.")

    with graph_tab:
        _render_section_intro(
            "Relationship Graph",
            "Use the left control rail to narrow the network, while keeping the graph centered for easier scanning.",
            kicker="Graph",
        )
        available_graph_weeks = (
            list_available_graph_weeks_neo4j(0)
            if data_backend == "Neo4j"
            else list_available_graph_weeks()
        )
        if not available_graph_weeks:
            if data_backend == "Neo4j":
                st.info("No Neo4j graph weeks found. Run the pipeline or sync Neo4j first.")
            else:
                st.info("Graph export not found. Run python -m src.export_graph --format json first.")
        else:
            latest_index = len(available_graph_weeks) - 1
            control_top_left, control_top_right = st.columns([0.55, 0.45], gap="large")
            with control_top_left:
                selected_graph_week = st.selectbox(
                    "Select Week",
                    options=available_graph_weeks,
                    index=latest_index,
                    key="graph_week_select",
                )
            selected_week_index = available_graph_weeks.index(selected_graph_week)
            previous_graph_week = (
                available_graph_weeks[selected_week_index - 1] if selected_week_index > 0 else None
            )
            with control_top_right:
                compare_with_previous = st.toggle(
                    "Compare with previous week",
                    value=False,
                    disabled=previous_graph_week is None,
                    key="graph_compare_toggle",
                )
                if compare_with_previous and previous_graph_week is None:
                    st.caption("No previous week snapshot is available yet.")

            if data_backend == "Neo4j":
                current_graph = load_graph_snapshot_neo4j(selected_graph_week, 0)
            else:
                current_graph_path = _graph_snapshot_path(selected_graph_week)
                current_graph = load_graph_snapshot(
                    selected_graph_week,
                    _path_version(current_graph_path),
                )
            if current_graph is None:
                st.info(f"Graph snapshot not found for {selected_graph_week}.")
                return

            previous_graph = None
            if compare_with_previous and previous_graph_week is not None:
                if data_backend == "Neo4j":
                    previous_graph = load_graph_snapshot_neo4j(previous_graph_week, 0)
                else:
                    previous_graph_path = _graph_snapshot_path(previous_graph_week)
                    previous_graph = load_graph_snapshot(
                        previous_graph_week,
                        _path_version(previous_graph_path),
                    )

            if compare_with_previous and previous_graph is not None:
                graph_payload, comparison_summary, emerging_entities = build_comparison_payload(
                    current_graph,
                    previous_graph,
                )
            else:
                graph_payload = current_graph
                comparison_summary = {
                    "current_nodes": len(current_graph.get("nodes", [])),
                    "previous_nodes": len(previous_graph.get("nodes", [])) if previous_graph else 0,
                    "new_nodes": 0,
                    "removed_nodes": 0,
                    "new_edges": 0,
                    "removed_edges": 0,
                }
                emerging_entities = []

            graph_node_types = sorted({str(node.get("type", "")) for node in graph_payload.get("nodes", [])})
            graph_relationships = sorted(
                {str(edge.get("relationship", "")) for edge in graph_payload.get("edges", [])}
            )
            with st.expander("Graph filters", expanded=False):
                _render_graph_legend()
                filter_cols = st.columns([0.24, 0.46, 0.30], gap="large")
                with filter_cols[0]:
                    selected_node_types = st.multiselect(
                        "Node types",
                        options=graph_node_types,
                        default=graph_node_types,
                        key="graph_node_types",
                    )
                with filter_cols[1]:
                    selected_relationship_types = st.multiselect(
                        "Relationship types",
                        options=graph_relationships,
                        default=graph_relationships,
                        key="graph_relationship_types",
                    )
                with filter_cols[2]:
                    graph_search = st.text_input("Search node", value="", key="graph_search")

            st.markdown(
                """
                <div class="graph-shell">
                  <div class="section-kicker">Canvas</div>
                  <div class="graph-stage-title">Expanded Network View</div>
                  <div class="graph-stage-copy">Hover nodes for importance and hover edges for a one-sentence narrative explanation.</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            render_graph(
                graph_payload,
                selected_node_types,
                selected_relationship_types,
                graph_search,
            )

            st.markdown(
                """
                <div class="graph-control-card">
                  <div class="section-kicker">Snapshot</div>
                  <div class="graph-control-title">Graph Summary</div>
                  <div class="graph-control-copy">Counts for the selected graph snapshot and optional comparison.</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            summary_cols = st.columns(6)
            with summary_cols[0]:
                _render_metric_card("Nodes This Week", str(comparison_summary["current_nodes"]), selected_graph_week)
            with summary_cols[1]:
                previous_label = previous_graph_week if previous_graph_week else "No prior snapshot"
                _render_metric_card("Nodes Last Week", str(comparison_summary["previous_nodes"]), previous_label)
            with summary_cols[2]:
                _render_metric_card("New Nodes", str(comparison_summary["new_nodes"]), "Added vs previous week")
            with summary_cols[3]:
                _render_metric_card("Removed Nodes", str(comparison_summary["removed_nodes"]), "Missing this week")
            with summary_cols[4]:
                _render_metric_card("New Edges", str(comparison_summary["new_edges"]), "New relationships")
            with summary_cols[5]:
                _render_metric_card("Removed Edges", str(comparison_summary["removed_edges"]), "Dropped relationships")

            if compare_with_previous and previous_graph_week is not None:
                if emerging_entities:
                    st.caption(
                        "Top emerging entities: "
                        + ", ".join(f"{label} (+{delta})" for label, delta in emerging_entities)
                    )
                else:
                    st.caption(f"Comparing {selected_graph_week} against {previous_graph_week}. No emerging entities yet.")


if __name__ == "__main__":
    main()
