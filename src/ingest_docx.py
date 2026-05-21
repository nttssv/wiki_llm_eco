"""DOCX ingestion helpers."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document

from .date_utils import parse_article_date


METADATA_ALIASES: dict[str, set[str]] = {
    "source": {"source", "publication", "publisher", "outlet"},
    "section": {"section", "desk"},
    "title": {"title", "headline"},
    "subtitle": {"subtitle", "subheading", "deck"},
    "published_date": {"date", "published", "published date", "publication date"},
    "category": {"category", "topic"},
}


DATE_LINE_PATTERN = re.compile(r"^\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}$")
INLINE_DATE_PATTERN = re.compile(
    r"\b("
    r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}"
    r"|"
    r"[A-Za-z]{3,9}\s+\d{1,2}(?:st|nd|rd|th)?\s+\d{4}"
    r")\b"
)
FILENAME_DATE_PATTERNS = (
    re.compile(r"(?P<date>\d{1,2}[A-Za-z]{3,9}\d{4})"),
    re.compile(r"(?P<date>\d{1,2}_[A-Za-z]{3,9}_\d{4})", re.IGNORECASE),
)
FILENAME_SOURCE_ALIASES = {
    "FT": "Financial Times",
    "WSJ": "Wall Street Journal",
    "NYT": "New York Times",
}
IGNORED_SUBTITLE_PARAGRAPHS = {
    "share",
    "listen to this story",
}
IGNORED_SUBTITLE_PREFIXES = (
    "illustration:",
    "photograph:",
)


@dataclass(slots=True)
class RawArticleDocument:
    """Parsed DOCX content used for extraction."""

    path: Path
    metadata: dict[str, str]
    body_text: str
    chart_references: list[str]
    paragraphs: list[str]


def _normalise_metadata_key(label: str) -> str | None:
    candidate = label.strip().lower()
    for canonical_key, aliases in METADATA_ALIASES.items():
        if candidate in aliases:
            return canonical_key
    return None


def _normalize_source_value(value: str) -> str:
    cleaned = value.strip()
    if not cleaned:
        return cleaned
    return FILENAME_SOURCE_ALIASES.get(cleaned.upper(), cleaned)


def _extract_metadata(paragraphs: list[str]) -> tuple[dict[str, str], set[int]]:
    metadata: dict[str, str] = {}
    consumed_indexes: set[int] = set()

    for index, paragraph in enumerate(paragraphs[:12]):
        match = re.match(
            r"^(?P<label>[A-Za-z][A-Za-z /-]{0,40}?)(?::\s*|\s+)(?P<value>.+)$",
            paragraph,
        )
        if not match:
            continue

        key = _normalise_metadata_key(match.group("label"))
        if key is None:
            continue

        value = match.group("value").strip()
        if key == "source":
            value = _normalize_source_value(value)
        metadata[key] = value
        consumed_indexes.add(index)

    return metadata, consumed_indexes


def _extract_filename_date(path: Path) -> str:
    stem = path.stem
    for pattern in FILENAME_DATE_PATTERNS:
        match = pattern.search(stem)
        if match is None:
            continue
        raw_date = match.group("date")
        if "_" in raw_date:
            candidate = raw_date.replace("_", " ").title()
        else:
            candidate = re.sub(r"^(\d{1,2})([A-Za-z]{3,9})(\d{4})$", r"\1 \2 \3", raw_date).title()
        parsed = parse_article_date(candidate)
        if parsed is not None:
            return parsed.strftime("%d %b %Y")
    return ""


def _extract_filename_source(path: Path) -> str:
    stem = path.stem
    last_token = stem.rsplit("_", 1)[-1].strip()
    if not last_token or not last_token.isalpha():
        return ""
    return _normalize_source_value(last_token)


def _extract_published_date(paragraphs: list[str], consumed_indexes: set[int]) -> str:
    for index, paragraph in enumerate(paragraphs[1:8], start=1):
        if index in consumed_indexes:
            continue

        if DATE_LINE_PATTERN.match(paragraph):
            parsed = parse_article_date(paragraph)
            if parsed is not None:
                consumed_indexes.add(index)
                return parsed.strftime("%d %b %Y")

        match = INLINE_DATE_PATTERN.search(paragraph)
        if match is None:
            continue

        parsed = parse_article_date(match.group(1))
        if parsed is None:
            continue

        consumed_indexes.add(index)
        return parsed.strftime("%d %b %Y")

    return ""


def _is_boilerplate_paragraph(paragraph: str) -> bool:
    cleaned = paragraph.strip().lower()
    return cleaned in IGNORED_SUBTITLE_PARAGRAPHS or cleaned.startswith(IGNORED_SUBTITLE_PREFIXES)


def read_docx_article(path: Path) -> RawArticleDocument:
    """Read a DOCX article and return lightweight parsed content."""

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    metadata, consumed_indexes = _extract_metadata(paragraphs)

    if paragraphs and "title" not in metadata:
        metadata["title"] = paragraphs[0]
        consumed_indexes.add(0)

    if "title" in metadata and "|" in metadata["title"] and not metadata.get("section"):
        section, title = (part.strip() for part in metadata["title"].split("|", 1))
        metadata["section"] = section
        metadata["title"] = title

    if "published_date" not in metadata:
        published_date = _extract_published_date(paragraphs, consumed_indexes)
        if published_date:
            metadata["published_date"] = published_date

    if not metadata.get("published_date"):
        published_date = _extract_filename_date(path)
        if published_date:
            metadata["published_date"] = published_date

    if metadata.get("source", "Unknown") == "Unknown":
        source = _extract_filename_source(path)
        if source:
            metadata["source"] = source

    if "subtitle" not in metadata:
        subtitle_parts: list[str] = []
        for index, paragraph in enumerate(paragraphs[1:5], start=1):
            if index in consumed_indexes:
                continue
            if _is_boilerplate_paragraph(paragraph):
                consumed_indexes.add(index)
                continue
            if DATE_LINE_PATTERN.match(paragraph):
                break
            if len(paragraph) > 180 or ":" in paragraph:
                break
            subtitle_parts.append(paragraph)
            consumed_indexes.add(index)

        if subtitle_parts:
            metadata["subtitle"] = " ".join(subtitle_parts)

    metadata.setdefault("source", "Unknown")
    metadata.setdefault("section", "")
    metadata.setdefault("subtitle", "")
    metadata.setdefault("published_date", "")
    metadata.setdefault("category", metadata["section"])

    body_paragraphs = [
        text
        for index, text in enumerate(paragraphs)
        if index not in consumed_indexes and not _is_boilerplate_paragraph(text)
    ]
    chart_references = [
        paragraph
        for paragraph in paragraphs
        if re.search(r"\b(chart|figure|graphic|image|exhibit)\b", paragraph, re.IGNORECASE)
    ]

    return RawArticleDocument(
        path=path,
        metadata=metadata,
        body_text="\n\n".join(body_paragraphs),
        chart_references=chart_references,
        paragraphs=paragraphs,
    )
